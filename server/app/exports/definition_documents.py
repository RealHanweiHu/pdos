"""Internal definition template: explicit fields, complete text, no remote content."""

import json
from io import BytesIO
from typing import Any, Literal
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt as SlidePt
from reportlab.lib import colors  # type: ignore[import-untyped]
from reportlab.lib.enums import TA_LEFT  # type: ignore[import-untyped]
from reportlab.lib.styles import ParagraphStyle  # type: ignore[import-untyped]
from reportlab.lib.units import mm  # type: ignore[import-untyped]
from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
from reportlab.pdfbase.cidfonts import UnicodeCIDFont  # type: ignore[import-untyped]
from reportlab.platypus import (  # type: ignore[import-untyped]
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

Format = Literal["docx", "xlsx", "pptx", "pdf"]
VERSION = "internal-definition.v1"
MIMES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
}


def value(v: Any) -> str:
    if v is None:
        return "未记录"
    if isinstance(v, str):
        return v
    return json.dumps(v, ensure_ascii=False, indent=2)


def sections(snapshot: dict[str, Any]) -> list[tuple[str, str]]:
    d, content = snapshot["definition"], snapshot["definition"]["content"]
    data = [
        (
            "适用范围",
            "内部资料。固定生成时状态，以系统当前版本为准。导出不代表客户批准、审厂或认证通过。",
        ),
        (
            "版本与状态",
            f"{snapshot['status_label']}\n定义 {d['id']}\n"
            f"修订 {d['revision_number']} · 内容版本 {d['version']}\n"
            f"发布时间 {d['published_at'] or '未发布'}\n模板 {VERSION}",
        ),
        ("对象范围", value(snapshot["product"])),
        (
            "缺项与来源提示",
            "\n".join(snapshot["warnings"])
            or "本次读取未发现已登记的来源提示；不构成资料完整性认证。",
        ),
    ]
    labels = {
        "target_users": "目标用户",
        "scenarios": "使用场景",
        "included_features": "包含功能",
        "excluded_features": "排除功能",
        "region_constraints": "地域约束",
        "target_cost": "成本目标（声明）",
        "price_position": "价格定位（声明）",
        "launch_window": "上市窗口（声明）",
        "service_requirements": "服务要求",
        "success_criteria": "成功判据",
        "change_reason": "修订原因",
    }
    for key, label in labels.items():
        data.append((label, value(content[key])))
    labels = {
        "state": "确认状态",
        "target_value": "数值目标",
        "target_text": "文字目标",
        "unit": "单位",
        "tolerance": "容差",
        "conditions": "工况",
        "priority": "优先级",
        "required_for_release": "发布必需",
        "verification_method": "验证方法",
        "acceptance_criterion": "接受判据",
        "verification_stage": "验证阶段",
        "owner_user_id": "负责人",
        "rationale": "依据与假设",
        "disposition_reason": "不适用说明",
    }
    for item in content["items"]:
        data.append(
            (
                f"规格 {item['item_key']} · {item['title']}",
                "\n\n".join(label + "：" + value(item.get(key)) for key, label in labels.items()),
            )
        )
    data.extend(
        [
            ("专业协审历史", value(snapshot["reviews"])),
            ("产品发布与撤回历史", value(snapshot["releases"])),
            ("原批准引用", value(d["approval"])),
            ("固定来源引用（未附原文）", value(snapshot["sources"])),
        ]
    )
    return data


def chunks(text: str, size: int) -> list[str]:
    """Bound document cells/slides without truncating any source characters."""
    return [text[i : i + size] for i in range(0, len(text), size)] or [""]


def render(snapshot: dict[str, Any], format: Format, receipt: dict[str, str]) -> bytes:
    return render_sections(
        snapshot["definition"]["title"],
        snapshot["status_label"],
        sections(snapshot),
        format,
        receipt,
    )


def render_sections(
    title: str,
    status: str,
    content: list[tuple[str, str]],
    format: Format,
    receipt: dict[str, str],
    *,
    audience: Literal["internal", "customer"] = "internal",
) -> bytes:
    audience_label = "客户" if audience == "customer" else "内部"
    body = [("导出回执", "\n".join(f"{k}：{v}" for k, v in receipt.items())), *content]
    stream = BytesIO()
    if format == "docx":
        document = Document()
        style = document.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(10)
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Cm(2)
        section.header.paragraphs[0].text = status + " · " + audience_label + "资料"
        section.footer.paragraphs[0].text = "固定生成时状态 · 回查编号 " + receipt["编号"]
        document.add_heading(title, 0)
        for heading, text in body:
            document.add_heading(heading, 1)
            for part in chunks(text, 4000):
                document.add_paragraph(part)
        document.save(stream)
    elif format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = audience_label + "报告"
        sheet.append(["章节 / 字段", "原值 / 固定记录"])
        for heading, text in [(title, status), *body]:
            for index, part in enumerate(chunks(text, 2000)):
                sheet.append([heading if index == 0 else heading + f"（续 {index + 1}）", part])
                for cell in sheet[sheet.max_row]:
                    # User-controlled strings never become formula or hyperlink cells.
                    cell.data_type = "s"
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 110
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="C2410C")
            cell.font = Font(color="FFFFFF", bold=True)
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        workbook.save(stream)
    elif format == "pptx":
        presentation = Presentation()
        presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
        for heading, text in [(title, status), *body]:
            if len(heading) > 36:
                text = heading + "\n\n" + text
                heading = heading[:36] + "…"
            # Every line and character is included, with explicit continuation slides.
            wrapped = [
                line[i : i + 54]
                for line in text.split("\n")
                for i in range(0, max(1, len(line)), 54)
            ]
            pages = [wrapped[i : i + 15] for i in range(0, len(wrapped), 15)] or [[""]]
            for index, lines in enumerate(pages):
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                title_box = slide.shapes.add_textbox(
                    Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.85)
                )
                title_box.text_frame.paragraphs[0].text = heading + (
                    f"（续 {index + 1}）" if index else ""
                )
                title_box.text_frame.paragraphs[0].font.size = SlidePt(24)
                box = slide.shapes.add_textbox(Inches(0.5), Inches(1.15), Inches(12.3), Inches(5.6))
                box.text_frame.word_wrap = True
                for i, line in enumerate(lines):
                    paragraph = (
                        box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
                    )
                    paragraph.text = line
                    paragraph.font.name = "Microsoft YaHei"
                    paragraph.font.size = SlidePt(16)
                    paragraph.space_after = SlidePt(3)
                footer = slide.shapes.add_textbox(Inches(0.5), Inches(7), Inches(12.3), Inches(0.3))
                footer.text_frame.paragraphs[0].text = (
                    status
                    + " · "
                    + audience_label
                    + " · "
                    + receipt["编号"]
                    + f" · {len(presentation.slides)}"
                )
                footer.text_frame.paragraphs[0].font.size = SlidePt(9)
        presentation.save(stream)
    else:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        normal = ParagraphStyle(
            "body",
            fontName="STSong-Light",
            fontSize=10,
            leading=15,
            wordWrap="CJK",
            alignment=TA_LEFT,
        )
        heading_style = ParagraphStyle(
            "heading",
            parent=normal,
            fontSize=16,
            leading=22,
            textColor=colors.HexColor("#c2410c"),
            spaceBefore=12,
            spaceAfter=6,
            keepWithNext=True,
        )
        story = [Paragraph(escape(title), heading_style)]
        for heading, text in body:
            story.append(Paragraph(escape(heading), heading_style))
            for part in chunks(text, 2000):
                story.append(Paragraph(escape(part).replace("\n", "<br/>"), normal))
            story.append(Spacer(1, 4 * mm))

        def footer(canvas: Any, document: Any) -> None:
            canvas.saveState()
            canvas.setFont("STSong-Light", 8)
            canvas.drawString(
                18 * mm,
                12 * mm,
                status + " · " + audience_label + " · " + receipt["编号"] + f" · {document.page}",
            )
            canvas.restoreState()

        SimpleDocTemplate(
            stream,
            pagesize=(210 * mm, 297 * mm),
            rightMargin=18 * mm,
            leftMargin=18 * mm,
            topMargin=18 * mm,
            bottomMargin=22 * mm,
            title=title,
            author="PDOS",
        ).build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()
