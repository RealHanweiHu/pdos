"""Human-confirmed XLSX supply captures with immutable source coordinates."""

from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any, Literal, Self
from uuid import UUID, uuid4

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from pydantic import AwareDatetime, Field, ValidationError, model_validator
from sqlalchemy import Connection, Engine, select

from app.auth.repository import AuthContext
from app.auth.roles import Action
from app.contracts.common_objects import (
    Contract,
    Digest,
    Key,
    Quantity,
    ScopeRef,
    SourceRef,
    SupplySnapshot,
    Text,
    Unit,
)
from app.contracts.time_operations import SupplySnapshotCapturedPayload
from app.db.schema import tables
from app.domain.document_previews import DocumentPreviewService, document_row, metadata
from app.domain.documents import _validate_xlsx_archive
from app.domain.knowledge_corrections import KnowledgeCorrectionService
from app.domain.material_supply_pools import lock_pools
from app.domain.product_workspace import _plain, digest
from app.domain.project_write_lock import lock_project_write_scope
from app.events.writer import AuditDraft, EventAuditWriter, EventDraft
from app.storage.documents import DocumentObjectStore

VERSION = "supply.xlsx.capture.v1"
MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FIELDS = {
    "external_pool_id": "供给池/库位批号",
    "external_item_id": "ERP物料标识",
    "quantity": "时点可分配量",
    "unit": "单位",
    "supply_kind": "现货/预计到货",
    "quality_status": "质量状态",
    "expected_date": "预计到货日期",
    "lead_time_days": "交期天数",
}
REQUIRED = set(FIELDS) - {"expected_date", "lead_time_days"}


class SupplyImportError(ValueError):
    pass


class InspectSupply(Contract):
    document_id: UUID


class SupplyImportDraft(InspectSupply):
    source_system: Key
    source_version: str = Field(min_length=1, max_length=256)
    coverage: Text
    authorization_basis: Text
    as_of: AwareDatetime
    fresh_until: AwareDatetime
    sheet: str = Field(min_length=1, max_length=128)
    header_row: int = Field(default=1, ge=1, le=100)
    columns: dict[str, str] = Field(min_length=6, max_length=8)

    @model_validator(mode="after")
    def mapping(self) -> Self:
        if self.fresh_until <= self.as_of:
            raise ValueError("有效期须晚于来源数据时点")
        if not set(self.columns) >= REQUIRED or not set(self.columns) <= set(FIELDS):
            raise ValueError("请完整映射必需字段，不支持额外字段")
        numbers = []
        for column in self.columns.values():
            try:
                number = column_index_from_string(column)
            except ValueError as error:
                raise ValueError("请选择实际工作表列") from error
            if not column.isupper() or not 1 <= number <= 100:
                raise ValueError("本次导入支持前100列")
            numbers.append(number)
        if len(set(numbers)) != len(numbers):
            raise ValueError("每个来源列只能映射一个字段")
        return self


class SupplyImportConfirm(SupplyImportDraft):
    request_id: UUID
    expected_digest: Digest
    reason: Text


class SupplyLine(Contract):
    external_pool_id: str = Field(min_length=1, max_length=256)
    external_item_id: str = Field(min_length=1, max_length=256)
    quantity: Quantity | None = Field(ge=0)
    unit: Unit
    supply_kind: Literal["on_hand", "expected_receipt"]
    quality_status: Literal["released", "frozen", "unknown"]
    expected_date: date | None = None
    lead_time_days: int | None = Field(default=None, ge=0, le=36600, strict=True)

    @model_validator(mode="after")
    def arrival(self) -> Self:
        if self.supply_kind == "on_hand" and self.expected_date is not None:
            raise ValueError("现货行不填预计到货日期；分开记录预计到货")
        return self


def _authorize(
    c: Connection, context: AuthContext, project_id: UUID, *, write: bool = False
) -> None:
    KnowledgeCorrectionService._authorize(
        c,
        context.tenant.id,
        context.user.id,
        project_id,
        Action.PRODUCT_ENGINEERING_EDIT if write else Action.PRODUCT_ENGINEERING_VIEW,
    )


def _read(
    engine: Engine,
    store: DocumentObjectStore,
    context: AuthContext,
    project_id: UUID,
    document_id: UUID,
) -> tuple[dict[str, Any], bytes]:
    with engine.connect() as c:
        _authorize(c, context, project_id)
        KnowledgeCorrectionService._authorize(
            c, context.tenant.id, context.user.id, project_id, Action.DOCUMENT_ANALYZE
        )
        row = document_row(c, context.tenant.id, project_id, document_id)
        if row["media_type"] != MEDIA:
            raise SupplyImportError("请选择本项目已保存的XLSX原件")
    meta, content, _ = DocumentPreviewService(engine, store).original(
        context.tenant.id, project_id, document_id
    )
    _validate_xlsx_archive(content)
    with engine.connect() as c:
        _authorize(c, context, project_id)
    return meta, content


def _rows(content: bytes, command: SupplyImportDraft) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if command.sheet not in workbook.sheetnames:
            raise SupplyImportError("原工作表不存在，请重新选择")
        sheet = workbook[command.sheet]
        if sheet.max_row is not None and sheet.max_row > command.header_row + 2000:
            raise SupplyImportError("本次最多2000个数据行，请按明确导出范围拆分原件")
        positions = {
            field: column_index_from_string(column) - 1 for field, column in command.columns.items()
        }
        lines, errors, ignored = [], [], []
        identities: set[tuple[str, str]] = set()
        for number, cells in enumerate(
            sheet.iter_rows(
                min_row=command.header_row + 1,
                max_row=command.header_row + 2001,
                max_col=max(positions.values()) + 1,
            ),
            command.header_row + 1,
        ):
            raw = {
                field: {
                    "cell": f"{command.columns[field]}{number}",
                    "value": cell.value.isoformat()
                    if isinstance(cell.value, (date, datetime))
                    else cell.value,
                    "type": cell.data_type,
                }
                for field, position in positions.items()
                for cell in [cells[position]]
            }
            if all(value["value"] in (None, "") for value in raw.values()):
                if sheet.max_row and number <= sheet.max_row:
                    ignored.append(number)
                continue
            if number > command.header_row + 2000:
                raise SupplyImportError("数据行超过本次2000行范围")
            try:
                if any(value["type"] in {"f", "e"} for value in raw.values()):
                    raise ValueError("映射列包含公式或错误值，请使用保留原出处的数值导出版本")
                values: dict[str, Any] = {field: value["value"] for field, value in raw.items()}
                for field in ("quantity", "expected_date", "lead_time_days"):
                    if values.get(field) in (None, ""):
                        values[field] = None
                values["supply_kind"] = {"现货": "on_hand", "预计到货": "expected_receipt"}.get(
                    str(values["supply_kind"]), values["supply_kind"]
                )
                values["quality_status"] = {
                    "已放行": "released",
                    "冻结": "frozen",
                    "未知": "unknown",
                }.get(str(values["quality_status"]), values["quality_status"])
                arrival = values.get("expected_date")
                if arrival is not None:
                    if not isinstance(arrival, str):
                        raise ValueError(
                            "预计到货日期须为日期单元格或ISO日期，不将普通数字当作日期"
                        )
                    arrival = arrival.strip()
                    values["expected_date"] = (
                        date.fromisoformat(arrival)
                        if len(arrival) == 10
                        else datetime.fromisoformat(arrival).date()
                    )
                if values.get("quality_status") in (None, ""):
                    values["quality_status"] = "unknown"
                if isinstance(values.get("quantity"), bool):
                    raise ValueError("数量不能使用是非值")
                if (
                    isinstance(values.get("lead_time_days"), str)
                    and values["lead_time_days"].strip().isdigit()
                ):
                    values["lead_time_days"] = int(values["lead_time_days"].strip())
                line = SupplyLine.model_validate(values)
                identity = (line.external_pool_id, line.external_item_id)
                if identity in identities:
                    raise ValueError(
                        "同一供给池和物料出现重复行；请在原导出中明确不同池或批号，不能重复计量"
                    )
                identities.add(identity)
                lines.append(
                    {"row": number, "values": line.model_dump(mode="json"), "raw_cells": raw}
                )
            except (ValidationError, ValueError) as error:
                errors.append({"row": number, "message": str(error), "raw_cells": raw})
        if not lines and not errors:
            raise SupplyImportError("所选范围没有供给数据")
        return {"lines": lines, "errors": errors, "ignored_rows": ignored}
    finally:
        workbook.close()


class SupplyImportService:
    def __init__(self, engine: Engine, store: DocumentObjectStore):
        self.engine, self.store = engine, store

    def list(self, context: AuthContext, project_id: UUID) -> dict[str, Any]:
        with self.engine.connect() as c:
            _authorize(c, context, project_id)
            batches = list(
                c.execute(
                    select(tables.supply_import_batches)
                    .where(
                        tables.supply_import_batches.c.tenant_id == context.tenant.id,
                        tables.supply_import_batches.c.project_id == project_id,
                    )
                    .order_by(
                        tables.supply_import_batches.c.captured_at.desc(),
                        tables.supply_import_batches.c.id,
                    )
                    .limit(51)
                ).mappings()
            )
            documents = list(
                c.execute(
                    select(
                        tables.documents.c.id,
                        tables.documents.c.filename,
                        tables.documents.c.sha256,
                    )
                    .where(
                        tables.documents.c.tenant_id == context.tenant.id,
                        tables.documents.c.project_id == project_id,
                        tables.documents.c.media_type == MEDIA,
                    )
                    .order_by(tables.documents.c.created_at.desc(), tables.documents.c.id)
                    .limit(201)
                ).mappings()
            )
            return _plain(
                {
                    "batches": [dict(row) for row in batches[:50]],
                    "has_more": len(batches) > 50,
                    "documents": [dict(row) for row in documents[:200]],
                    "documents_has_more": len(documents) > 200,
                    "fields": FIELDS,
                    "now": datetime.now(UTC),
                }
            )

    def inspect(self, context: AuthContext, project_id: UUID, document_id: UUID) -> dict[str, Any]:
        meta, content = _read(self.engine, self.store, context, project_id, document_id)
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            return {
                "document": meta,
                "sheets": [
                    {
                        "name": sheet.title,
                        "headers": [
                            {
                                "column": get_column_letter(index),
                                "label": str(cell.value or "")[:256],
                            }
                            for index, cell in enumerate(
                                next(
                                    sheet.iter_rows(
                                        min_row=1,
                                        max_row=1,
                                        max_col=min(sheet.max_column or 1, 100),
                                    ),
                                    (),
                                ),
                                1,
                            )
                        ],
                    }
                    for sheet in workbook.worksheets[:50]
                ],
            }
        finally:
            workbook.close()

    def preview(
        self, context: AuthContext, project_id: UUID, command: SupplyImportDraft
    ) -> dict[str, Any]:
        if command.as_of > datetime.now(UTC):
            raise SupplyImportError("来源数据时点不能是未来")
        meta, content = _read(self.engine, self.store, context, project_id, command.document_id)
        parsed = _rows(content, command)
        configuration = SupplyImportDraft.model_validate(
            command.model_dump(exclude={"request_id", "expected_digest", "reason"})
        ).model_dump(mode="json")
        return {
            "document": meta,
            "configuration": configuration,
            **parsed,
            "digest": digest({"document": meta, "configuration": configuration, **parsed}),
            "stale": command.fresh_until <= datetime.now(UTC),
        }

    def detail(self, context: AuthContext, project_id: UUID, batch_id: UUID) -> dict[str, Any]:
        with self.engine.connect() as c:
            _authorize(c, context, project_id)
            batch = (
                c.execute(
                    select(tables.supply_import_batches).where(
                        tables.supply_import_batches.c.tenant_id == context.tenant.id,
                        tables.supply_import_batches.c.project_id == project_id,
                        tables.supply_import_batches.c.id == batch_id,
                    )
                )
                .mappings()
                .one_or_none()
            )
            if batch is None:
                raise SupplyImportError("供给记录不存在于当前可读项目")
            rows = list(
                c.execute(
                    select(tables.supply_snapshots)
                    .where(
                        tables.supply_snapshots.c.tenant_id == context.tenant.id,
                        tables.supply_snapshots.c.batch_id == batch_id,
                    )
                    .order_by(tables.supply_snapshots.c.source_row)
                ).mappings()
            )
            return _plain(
                {
                    "batch": dict(batch),
                    "items": [dict(row) for row in rows],
                    "stale": batch["fresh_until"] <= datetime.now(UTC),
                    "source_href": f"/api/document-previews/projects/{project_id}/"
                    f"{batch['document_id']}/original",
                }
            )

    def confirm(
        self, context: AuthContext, project_id: UUID, command: SupplyImportConfirm
    ) -> dict[str, Any]:
        signature = digest(
            {"project_id": str(project_id), "command": command.model_dump(mode="json")}
        )
        # Check known receipts before re-reading the object store on a retry.
        with self.engine.connect() as c:
            _authorize(c, context, project_id, write=True)
            receipt = (
                c.execute(
                    select(tables.supply_import_batches).where(
                        tables.supply_import_batches.c.tenant_id == context.tenant.id,
                        tables.supply_import_batches.c.created_by_user_id == context.user.id,
                        tables.supply_import_batches.c.request_id == command.request_id,
                    )
                )
                .mappings()
                .one_or_none()
            )
            if receipt:
                if receipt["request_digest"] != signature:
                    raise SupplyImportError("请求编号已用于其他内容，请按原请求核对")
                return _plain(dict(receipt))
        preview = self.preview(context, project_id, command)
        if preview["errors"] or preview["digest"] != command.expected_digest:
            raise SupplyImportError("存在未处理的原行错误或预览已变化，请重新核对全部行")
        with self.engine.begin() as c:
            _authorize(c, context, project_id, write=True)
            lock_project_write_scope(c, tenant_id=context.tenant.id, project_id=project_id)
            current = metadata(document_row(c, context.tenant.id, project_id, command.document_id))
            if current != preview["document"]:
                raise SupplyImportError("原件登记已变化，请重新核对")
            batches = tables.supply_import_batches
            existing = (
                c.execute(
                    select(batches).where(
                        batches.c.tenant_id == context.tenant.id,
                        batches.c.source_system == command.source_system,
                        batches.c.source_version == command.source_version,
                    )
                )
                .mappings()
                .one_or_none()
            )
            if existing:
                if (
                    existing["project_id"] == project_id
                    and existing["content_digest"] == preview["digest"]
                ):
                    return _plain(dict(existing))
                raise SupplyImportError(
                    "同一来源的导出版本已登记且内容或项目不同，请核对已有授权记录，不能重复建供给池"
                )
            lock_pools(
                c,
                context.tenant.id,
                [
                    (
                        command.source_system,
                        row["values"]["external_pool_id"],
                        row["values"]["external_item_id"],
                    )
                    for row in preview["lines"]
                ],
            )
            batch_id, captured_at = uuid4(), datetime.now(UTC)
            batch = (
                c.execute(
                    batches.insert()
                    .values(
                        id=batch_id,
                        tenant_id=context.tenant.id,
                        project_id=project_id,
                        document_id=command.document_id,
                        document_digest=current["sha256"],
                        source_system=command.source_system,
                        source_version=command.source_version,
                        as_of=command.as_of,
                        fresh_until=command.fresh_until,
                        captured_at=captured_at,
                        configuration=preview["configuration"],
                        content_digest=preview["digest"],
                        created_by_user_id=context.user.id,
                        request_id=command.request_id,
                        request_digest=signature,
                        reason=command.reason,
                        row_count=len(preview["lines"]),
                        ignored_rows=preview["ignored_rows"],
                    )
                    .returning(batches)
                )
                .mappings()
                .one()
            )
            writer = EventAuditWriter(
                c, tenant_id=context.tenant.id, project_id=project_id, actor_user_id=context.user.id
            )
            for row in preview["lines"]:
                line = SupplyLine.model_validate(row["values"])
                identity = uuid4()
                cells = ",".join(value["cell"] for value in row["raw_cells"].values())
                locator = f"{command.sheet}!row {row['row']} ({cells})"
                contract = SupplySnapshot(
                    id=identity,
                    tenant_id=context.tenant.id,
                    scope=ScopeRef(kind="project", ref_id=project_id),
                    version=1,
                    owner_user_id=context.user.id,
                    sources=(
                        SourceRef(
                            tenant_id=context.tenant.id,
                            project_id=project_id,
                            kind="document",
                            ref_id=command.document_id,
                            as_of=datetime.fromisoformat(current["created_at"]),
                            digest=current["sha256"],
                            locator=locator,
                        ),
                    ),
                    source_system=command.source_system,
                    external_item_id=line.external_item_id,
                    as_of=command.as_of,
                    fresh_until=command.fresh_until,
                    quantity=line.quantity,
                    unit=line.unit,
                    availability="unknown" if line.quantity is None else "known",
                )
                c.execute(
                    tables.supply_snapshots.insert().values(
                        id=identity,
                        tenant_id=context.tenant.id,
                        project_id=project_id,
                        batch_id=batch_id,
                        external_pool_id=line.external_pool_id,
                        external_item_id=line.external_item_id,
                        quantity=line.quantity,
                        unit=line.unit,
                        availability=contract.availability,
                        supply_kind=line.supply_kind,
                        quality_status=line.quality_status,
                        expected_date=line.expected_date,
                        lead_time_days=line.lead_time_days,
                        source_row=row["row"],
                        raw_cells=row["raw_cells"],
                        contract=contract.model_dump(mode="json"),
                    )
                )
                payload = SupplySnapshotCapturedPayload(
                    object_id=identity,
                    object_version=1,
                    operation_id=command.request_id,
                    correlation_id=batch_id,
                    causation_event_row_id=None,
                    source_id=batch_id,
                    snapshot_id=identity,
                    source_digest=current["sha256"],
                )
                writer.record(
                    EventDraft(
                        event_id=f"supply:{batch_id}:{identity}",
                        event_type="SUPPLY_SNAPSHOT_CAPTURED",
                        source="supply_snapshot_service",
                        entity_id=identity,
                        payload=payload.model_dump(mode="json"),
                    ),
                    AuditDraft(
                        action="supply.snapshot_captured",
                        after_snapshot=contract.model_dump(mode="json"),
                        context={
                            "batch_id": str(batch_id),
                            "source_row": row["row"],
                            "reason": command.reason,
                        },
                    ),
                )
            return _plain(dict(batch))
