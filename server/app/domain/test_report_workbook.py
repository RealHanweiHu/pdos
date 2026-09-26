"""Deterministic layout adapters; no product policy and no model calls."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

FORM_HEADERS = ["用例名称", "测试结果", "样机SN", "备注"]
_NON_DATA_SHEETS = ("封面", "目录", "汇总", "版本更新", "问题点")
_RECORD_STARTS = {"测试记录", "测试环境", "实测记录-图片", "测试前记录"}
_RECORD_LABELS = {
    "序号",
    "No",
    "图片",
    "备注",
    "实测记录",
    "测试记录",
    "测试项目",
    "测试步骤/方法",
    "测试按键",
    "记录项",
    "测试前记录",
    "测试后记录",
    "样机编号",
    "测试距离",
    "测试现象记录",
    "测试前样机状态图片",
    "测试后样机状态图片",
    "测试前样机状态描述",
    "测试后样机状态描述",
}


class ImportError_(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


def normalize(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _json_value(value: object) -> str | int | float | bool | None:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


@dataclass(frozen=True)
class ReportRow:
    values: dict[str, object]
    source: dict[str, Any]


@dataclass(frozen=True)
class ParsedWorkbook:
    layout: str
    headers: list[str]
    rows: list[ReportRow]
    skipped: list[dict[str, object]]


def _label_value(rows: list[tuple[int, list[object]]], label: str) -> tuple[object, str | None]:
    # These named forms anchor labels at A/D, G/J and M/P. Never search
    # arbitrarily to the right: a missing value must not consume the next label.
    for row_number, cells in rows:
        for index in (0, 6, 12):
            if index < len(cells) and normalize(cells[index]) == label:
                value_index = index + 3
                value = cells[value_index] if value_index < len(cells) else None
                return value, f"{get_column_letter(value_index + 1)}{row_number}"
    return None, None


_BUILD_LABELS = {"试制批次", "试制批次编号", "build code", "build_code"}


def _build_hints(
    rows: list[tuple[int, list[object]]], headers: list[str] | None = None
) -> list[dict[str, Any]]:
    hints = []
    for row_number, cells in rows:
        for index, label in enumerate(headers if headers is not None else cells):
            if headers is None and index not in (0, 6, 12):
                continue
            if normalize(label).casefold() not in _BUILD_LABELS:
                continue
            value_index = index if headers is not None else index + 3
            value = cells[value_index] if value_index < len(cells) else None
            hints.append(
                {
                    "label": normalize(label),
                    "raw_value": _json_value(value),
                    "code": normalize(value),
                    "cell": f"{get_column_letter(value_index + 1)}{row_number}",
                }
            )
    return hints


def _form_row(name: str, rows: list[tuple[int, list[object]]]) -> ReportRow | None:
    title, title_cell = _label_value(rows[:6], "测试项目")
    standard, _ = _label_value(rows[:8], "判定标准")
    if standard is None:
        standard, _ = _label_value(rows[:8], "测试标准")
    if not normalize(title) or standard is None:
        return None
    verdict, result_cell = _label_value(rows, "结论")
    reported_date, date_cell = _label_value(rows[:6], "测试日期")
    # Preserve day precision. The service's midnight storage representation
    # must not be interpreted as an observed execution time.
    day = (
        reported_date.date().isoformat()
        if isinstance(reported_date, datetime)
        else reported_date.isoformat()
        if isinstance(reported_date, date)
        else None
    )
    model, _ = _label_value(rows[:6], "样机型号")
    if model is None:
        model, _ = _label_value(rows[:6], "型号")
    method, _ = _label_value(rows[:8], "测试方法")
    if method is None:
        method, _ = _label_value(rows[:8], "检测方法")
    equipment, _ = _label_value(rows[:8], "测试仪器")
    if equipment is None:
        equipment, _ = _label_value(rows[:8], "检测仪器")
    records: list[dict[str, Any]] = []
    collecting = False
    has_observations = False
    observation_columns: set[int] = set()
    for row_number, cells in rows:
        first = normalize(cells[0]) if cells else ""
        if first == "结论":
            break
        if first in _RECORD_STARTS:
            collecting = True
            continue
        if not collecting:
            continue
        values = {
            f"{get_column_letter(i + 1)}{row_number}": _json_value(value)
            for i, value in enumerate(cells)
            if normalize(value)
        }
        records.append({"row": row_number, "cells": values})
        # Only explicit observation rows, never a table header or picture
        # caption, count as textual evidence. Other layouts remain reviewable.
        # Read only explicitly named observation columns. Methods and images
        # in adjacent merged cells cannot satisfy the evidence requirement.
        observation_labels = {
            "测试记录",
            "实测记录",
            "实测流量值",
            "实测自动关机时长",
            "实测噪音数据",
            "测试现象记录",
            "支持高压启停",
            "打气结束电压",
            "停机时显示气压",
            "重启后显示气压",
            "测试前样机状态描述",
            "测试后样机状态描述",
        }
        for column, value in enumerate(cells):
            if normalize(value) in observation_labels:
                if column == 0 and normalize(value).endswith("状态描述"):
                    observation_columns.add(3)
                elif column > 0:
                    observation_columns.add(column)
        is_header = first in {
            "序号",
            "测试项目",
            "记录项",
            "样机编号",
            "测试距离",
            "供电方式",
            "轮胎类型及尺寸",
        }
        if not is_header and any(
            column < len(cells)
            and normalize(cells[column])
            and normalize(cells[column]) not in observation_labels | _RECORD_LABELS
            for column in observation_columns
        ):
            has_observations = True
        observation_row = isinstance(cells[0], (int, float)) or first.endswith("状态描述")
        if observation_row and any(isinstance(value, (int, float)) for value in cells[3:]):
            has_observations = True
    # These are worksheet-level report observations, NOT inferred individual
    # sample runs. Serial number 1 in a table is intentionally not bound as SN.
    return ReportRow(
        values={"用例名称": name.strip(), "测试结果": normalize(verdict), "样机SN": "", "备注": ""},
        source={
            "sheet": name,
            "row": int(result_cell.lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
            if result_cell
            else rows[0][0],
            "build_hints": _build_hints(rows[:8]),
            "title": normalize(title),
            "title_cell": title_cell,
            "result_cell": result_cell,
            "reported_result": normalize(verdict),
            "model": normalize(model),
            "reported_date": day,
            "date_cell": date_cell,
            "raw_date": _json_value(reported_date),
            "method": normalize(method),
            "standard": normalize(standard),
            "equipment": normalize(equipment),
            "records": records,
            "has_observations": has_observations,
            "date_precision": "day" if day else None,
            "granularity": "worksheet_summary",
            "images_reviewed": False,
        },
    )


def parse_workbook(content: bytes) -> ParsedWorkbook:
    try:
        if len(content) > 20 * 1024 * 1024:
            raise ImportError_("IMPORT_WORKBOOK_TOO_LARGE", "workbook exceeds 20 MiB")
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ImportError_("IMPORT_WORKBOOK_TOO_LARGE", "expanded workbook exceeds limit")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except ImportError_:
        raise
    except Exception as error:  # noqa: BLE001 - untrusted workbook
        raise ImportError_("IMPORT_INVALID_XLSX", "spreadsheet could not be parsed") from error
    forms: list[ReportRow] = []
    flat_sheets: list[tuple[str, list[tuple[int, list[object]]]]] = []
    skipped: list[dict[str, object]] = []
    try:
        if len(workbook.worksheets) > 128:
            raise ImportError_("IMPORT_WORKBOOK_TOO_LARGE", "too many worksheets")
        for sheet in workbook.worksheets:
            if (sheet.max_row or 0) > 10_000 or (sheet.max_column or 0) > 128:
                raise ImportError_("IMPORT_WORKBOOK_TOO_LARGE", "worksheet dimensions exceed limit")
            rows: list[tuple[int, list[object]]] = [
                (index, list(row))
                for index, row in enumerate(sheet.iter_rows(values_only=True), 1)
                if any(normalize(value) for value in row)
            ]
            if not rows:
                continue
            if any(token in sheet.title for token in _NON_DATA_SHEETS):
                skipped.append({"sheet": sheet.title, "reason": "non_test_sheet"})
                continue
            form = _form_row(sheet.title, rows)
            if form is not None:
                forms.append(form)
            else:
                flat_sheets.append((sheet.title, rows))
    finally:
        workbook.close()
    if forms:
        skipped.extend(
            {"sheet": name, "reason": "unsupported_sheet_layout"} for name, _ in flat_sheets
        )
        return ParsedWorkbook("test_form_v1", FORM_HEADERS, forms, skipped)
    headers: list[str] = []
    flat_rows: list[ReportRow] = []
    for name, rows in flat_sheets:
        current = [normalize(value) for value in rows[0][1]]
        while current and not current[-1]:
            current.pop()
        if (
            len(current) < 2
            or any(not value for value in current)
            or len(set(current)) != len(current)
        ):
            skipped.append({"sheet": name, "reason": "unsupported_sheet_layout"})
            continue
        if headers and current != headers:
            skipped.append({"sheet": name, "reason": "different_flat_headers"})
            continue
        headers = current
        flat_rows.extend(
            ReportRow(
                {header: cells[i] if i < len(cells) else None for i, header in enumerate(headers)},
                {
                    "sheet": name,
                    "row": index,
                    "headers": headers,
                    "build_hints": _build_hints([(index, cells)], headers),
                    "cells": {
                        f"{get_column_letter(i + 1)}{index}": _json_value(value)
                        for i, value in enumerate(cells)
                    },
                },
            )
            for index, cells in rows[1:]
        )
    if not flat_rows:
        raise ImportError_("IMPORT_EMPTY_SHEET", "no supported test table or test form was found")
    return ParsedWorkbook("flat_v1", headers, flat_rows, skipped)
