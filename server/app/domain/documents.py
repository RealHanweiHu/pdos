"""Safe PDF/XLSX extraction and source-document persistence."""

from __future__ import annotations

import hashlib
import re
from contextlib import suppress
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from pathlib import PurePath
from uuid import UUID, uuid4
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from sqlalchemy import Engine

from app.db.schema import tables
from app.events.writer import AuditDraft, EventAuditWriter, EventDraft
from app.storage.documents import DocumentObjectStore

MAX_DOCUMENT_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_XLSX_BYTES = 100 * 1024 * 1024
MAX_DOCUMENT_CHARACTERS = 2_000_000
MAX_PDF_PAGES = 500
MAX_FRAGMENT_CHARACTERS = 20_000
SUPPORTED_MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class DocumentError(ValueError):
    code = "DOCUMENT_INVALID"


class UnsupportedDocumentTypeError(DocumentError):
    code = "DOCUMENT_TYPE_UNSUPPORTED"


class DocumentTooLargeError(DocumentError):
    code = "DOCUMENT_TOO_LARGE"


class DocumentExtractionError(DocumentError):
    code = "DOCUMENT_EXTRACTION_FAILED"


class EmptySpreadsheetError(DocumentError):
    code = "DOCUMENT_NO_TEXT"


@dataclass(frozen=True)
class ExtractedFragment:
    ordinal: int
    locator: str
    text: str


@dataclass(frozen=True)
class DocumentRecord:
    id: UUID
    project_id: UUID
    filename: str
    media_type: str
    classification: str
    size_bytes: int
    sha256: str
    extraction_status: str
    fragments: tuple[ExtractedFragment, ...]
    created_at: datetime


def _safe_filename(filename: str) -> str:
    normalized = PurePath(filename.replace("\\", "/")).name.strip()
    if not normalized or len(normalized) > 512:
        raise UnsupportedDocumentTypeError("document filename is invalid")
    return normalized


def _normalized_text(value: str) -> str:
    return "\n".join(line.rstrip() for line in value.replace("\x00", "").splitlines()).strip()


def _fragment(locator: str, text: str, *, start_ordinal: int) -> list[ExtractedFragment]:
    normalized = _normalized_text(text)
    if not normalized:
        return []
    chunks: list[ExtractedFragment] = []
    for offset in range(0, len(normalized), MAX_FRAGMENT_CHARACTERS):
        part = len(chunks) + 1
        part_locator = (
            locator if len(normalized) <= MAX_FRAGMENT_CHARACTERS else f"{locator}#part:{part}"
        )
        chunks.append(
            ExtractedFragment(
                ordinal=start_ordinal + len(chunks),
                locator=part_locator,
                text=normalized[offset : offset + MAX_FRAGMENT_CHARACTERS],
            )
        )
    return chunks


def _extract_pdf(content: bytes) -> tuple[ExtractedFragment, ...]:
    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise DocumentExtractionError("encrypted PDFs are not supported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentTooLargeError("PDF page count exceeds the v0 limit")
        fragments: list[ExtractedFragment] = []
        total_characters = 0
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            total_characters += len(text)
            if total_characters > MAX_DOCUMENT_CHARACTERS:
                raise DocumentTooLargeError("extracted PDF text exceeds the v0 limit")
            fragments.extend(
                _fragment(f"page:{page_number}", text, start_ordinal=len(fragments) + 1)
            )
        return tuple(fragments)
    except DocumentError:
        raise
    except (PdfReadError, ValueError) as error:
        raise DocumentExtractionError("PDF could not be parsed") from error


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            expanded = sum(item.file_size for item in archive.infolist())
            if expanded > MAX_EXPANDED_XLSX_BYTES:
                raise DocumentTooLargeError("expanded spreadsheet exceeds the v0 limit")
    except BadZipFile as error:
        raise DocumentExtractionError("spreadsheet is not a valid XLSX archive") from error


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"[\t\r\n]+", " ", str(value)).strip()


def _extract_xlsx(content: bytes) -> tuple[ExtractedFragment, ...]:
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, KeyError, ValueError) as error:
        raise DocumentExtractionError("spreadsheet could not be parsed") from error
    fragments: list[ExtractedFragment] = []
    total_characters = 0
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            batch_start: int | None = None
            batch_end: int | None = None
            batch_characters = 0
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                row_text = "\t".join(_cell_text(value) for value in values).rstrip()
                if not row_text:
                    continue
                if rows and batch_characters + len(row_text) + 1 > MAX_FRAGMENT_CHARACTERS:
                    locator = f"sheet:{sheet.title},rows:{batch_start}-{batch_end}"
                    fragments.extend(
                        _fragment(locator, "\n".join(rows), start_ordinal=len(fragments) + 1)
                    )
                    rows = []
                    batch_start = row_number
                    batch_characters = 0
                if batch_start is None:
                    batch_start = row_number
                rows.append(row_text)
                batch_end = row_number
                batch_characters += len(row_text) + 1
                total_characters += len(row_text) + 1
                if total_characters > MAX_DOCUMENT_CHARACTERS:
                    raise DocumentTooLargeError("spreadsheet text exceeds the v0 limit")
            if rows:
                locator = f"sheet:{sheet.title},rows:{batch_start}-{batch_end}"
                fragments.extend(
                    _fragment(locator, "\n".join(rows), start_ordinal=len(fragments) + 1)
                )
    finally:
        workbook.close()
    if not fragments:
        raise EmptySpreadsheetError("spreadsheet contains no extractable text")
    return tuple(fragments)


def extract_document(*, filename: str, content: bytes) -> tuple[str, tuple[ExtractedFragment, ...]]:
    safe_name = _safe_filename(filename)
    if not content:
        raise DocumentExtractionError("document is empty")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise DocumentTooLargeError("document exceeds the 20 MiB v0 limit")
    suffix = PurePath(safe_name).suffix.lower()
    media_type = SUPPORTED_MEDIA_TYPES.get(suffix)
    if media_type is None:
        raise UnsupportedDocumentTypeError("only PDF and XLSX are supported")
    if suffix == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise DocumentExtractionError("file content is not a PDF")
        return media_type, _extract_pdf(content)
    return media_type, _extract_xlsx(content)


class DocumentService:
    def __init__(self, engine: Engine, *, object_store: DocumentObjectStore) -> None:
        self._engine = engine
        self._object_store = object_store

    def analyze(
        self,
        *,
        tenant_id: UUID,
        project_id: UUID,
        actor_user_id: UUID,
        filename: str,
        classification: str,
        content: bytes,
        now: datetime | None = None,
    ) -> DocumentRecord:
        safe_name = _safe_filename(filename)
        media_type, fragments = extract_document(filename=safe_name, content=content)
        created_at = (now or datetime.now(UTC)).astimezone(UTC)
        document_id = uuid4()
        digest = hashlib.sha256(content).hexdigest()
        object_key = (
            f"tenants/{tenant_id}/projects/{project_id}/documents/{document_id}/{safe_name}"
        )
        extraction_status = "extracted" if fragments else "ocr_required"
        self._object_store.put(
            object_key=object_key,
            content=content,
            media_type=media_type,
            sha256=digest,
        )
        try:
            with self._engine.begin() as connection:
                connection.execute(
                    tables.documents.insert(),
                    {
                        "id": document_id,
                        "tenant_id": tenant_id,
                        "project_id": project_id,
                        "uploaded_by_user_id": actor_user_id,
                        "object_key": object_key,
                        "filename": safe_name,
                        "media_type": media_type,
                        "classification": classification,
                        "size_bytes": len(content),
                        "sha256": digest,
                        "extraction_status": extraction_status,
                        "fragment_count": len(fragments),
                        "created_at": created_at,
                    },
                )
                if fragments:
                    connection.execute(
                        tables.document_fragments.insert(),
                        [
                            {
                                "id": uuid4(),
                                "tenant_id": tenant_id,
                                "project_id": project_id,
                                "document_id": document_id,
                                "ordinal": item.ordinal,
                                "locator": item.locator,
                                "text": item.text,
                                "created_at": created_at,
                            }
                            for item in fragments
                        ],
                    )
                EventAuditWriter(
                    connection,
                    tenant_id=tenant_id,
                    actor_user_id=actor_user_id,
                    project_id=project_id,
                ).record(
                    EventDraft(
                        event_id=f"document:{document_id}:uploaded",
                        event_type="DOCUMENT_UPLOADED",
                        source="document_service",
                        entity_id=document_id,
                        payload={
                            "document_id": str(document_id),
                            "classification": classification,
                        },
                        occurred_at=created_at,
                    ),
                    AuditDraft(
                        action="document.upload",
                        after_snapshot={
                            "document_id": str(document_id),
                            "filename": safe_name,
                            "classification": classification,
                            "size_bytes": len(content),
                            "sha256": digest,
                            "extraction_status": extraction_status,
                            "fragment_count": len(fragments),
                        },
                    ),
                )
        except Exception:
            with suppress(Exception):
                self._object_store.delete(object_key=object_key)
            raise
        return DocumentRecord(
            id=document_id,
            project_id=project_id,
            filename=safe_name,
            media_type=media_type,
            classification=classification,
            size_bytes=len(content),
            sha256=digest,
            extraction_status=extraction_status,
            fragments=fragments,
            created_at=created_at,
        )
