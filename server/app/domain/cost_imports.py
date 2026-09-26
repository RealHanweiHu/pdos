"""Human-reconciled ERP XLSX snapshots atomically enter the actual cost ledger."""

from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any, Self
from uuid import UUID, uuid4, uuid5

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from pydantic import Field, model_validator
from sqlalchemy import Connection, Engine, func, select

from app.auth.repository import AuthContext
from app.contracts.common_objects import Contract, Digest, Key, Text
from app.db.schema import tables
from app.domain.build_materials import BuildMaterialService
from app.domain.cost_freezes import row
from app.domain.cost_inputs import CostInputService, InputApply, PriceDraft, current_query
from app.domain.cost_inputs import preview as input_preview
from app.domain.document_previews import document_row, metadata
from app.domain.npi_execution import NpiExecutionService
from app.domain.product_workspace import ProductWorkError, _plain, digest
from app.domain.supply_imports import MEDIA, SupplyImportService, _read
from app.storage.documents import DocumentObjectStore

ALGORITHM = "cost-actual.xlsx.v1"
FIELDS = {
    "source_line": "ERP唯一行标识",
    "title": "费用名称",
    "amount": "原币行总额",
    "quantity": "实际数量",
    "unit": "原单位",
    "currency": "币种",
    "subsystem": "子系统",
    "category": "费用类型",
    "basis": "单位／NRE",
    "tax_basis": "税费口径",
    "tax_rate": "税率（小数）",
    "freight_basis": "运费口径",
    "part_number": "ERP料号（人工对应）",
    "supplier_code": "ERP供方标识（人工对应）",
}
REQUIRED = {"source_line", "title", "amount", "quantity", "unit"}


class CostImportInspect(Contract):
    document_id: UUID


class CostImportAssignment(Contract):
    row: int = Field(ge=2, le=600)
    bom_item_id: UUID | None = None
    supplier_id: UUID | None = None
    mapping_basis: Text | None = None
    exclude_reason: Text | None = None

    @model_validator(mode="after")
    def consistent(self) -> Self:
        if self.exclude_reason and (self.bom_item_id or self.supplier_id or self.mapping_basis):
            raise ValueError("排除行只记录理由，不同时关联业务对象")
        if (self.bom_item_id or self.supplier_id) and not self.mapping_basis:
            raise ValueError("对应实际物料／供方时须具名说明核对依据")
        return self


class CostImportDraft(CostImportInspect):
    source_system: Key
    source_version: str = Field(min_length=1, max_length=120)
    as_of: date
    coverage: Text
    authorization_basis: Text
    reason: Text
    sheet: str = Field(min_length=1, max_length=128)
    header_row: int = Field(default=1, ge=1, le=100)
    columns: dict[str, str] = Field(min_length=5, max_length=14)
    defaults: dict[str, str] = Field(max_length=9)
    assignments: list[CostImportAssignment] = Field(default_factory=list, max_length=500)

    @model_validator(mode="after")
    def consistent(self) -> Self:
        if not self.columns.keys() >= REQUIRED or not self.columns.keys() <= FIELDS.keys():
            raise ValueError("请映射唯一行标识、费用名称、原币总额、实际数量和单位")
        allowed = set(FIELDS) - REQUIRED - {"part_number", "supplier_code"}
        if not self.defaults.keys() <= allowed or set(self.defaults) & set(self.columns):
            raise ValueError("同一字段只采用原列或统一口径之一")
        if (
            not {"currency", "subsystem", "category", "basis", "tax_basis", "freight_basis"}
            <= self.columns.keys() | self.defaults.keys()
        ):
            raise ValueError("请映射或明确填写币种、子系统、费用类型、单位/NRE及税费/运费口径")
        numbers = []
        for col in self.columns.values():
            number = column_index_from_string(col)
            if not col.isupper() or not 1 <= number <= 100:
                raise ValueError("只支持原工作表前100列")
            numbers.append(number)
        if len(set(numbers)) != len(numbers):
            raise ValueError("每个原列只能映射一个字段")
        if len({a.row for a in self.assignments}) != len(self.assignments):
            raise ValueError("同一原行不能重复对应")
        return self


class CostImportConfirm(Contract):
    request_id: UUID
    expected_digest: Digest
    draft: CostImportDraft
    evidence_checked: bool
    mapping_checked: bool
    exclusions_checked: bool


ALIASES = {
    "basis": {"单位": "unit", "NRE": "nre", "整项": "nre"},
    "category": {
        "物料": "material",
        "人工": "labor",
        "开模": "tooling",
        "工装": "tooling",
        "认证": "certification",
        "打样": "prototype",
        "测试": "testing",
        "物流": "logistics",
        "其他": "other",
    },
    "tax_basis": {
        "含税": "inclusive",
        "未含税": "exclusive",
        "不含税": "exclusive",
        "未知": "unknown",
    },
    "freight_basis": {
        "含运费": "included",
        "未含运费": "excluded",
        "不含运费": "excluded",
        "未知": "unknown",
    },
}


def cell_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, (Decimal, timedelta)):
        return str(value)
    return {
        "unsupported_type": type(value).__name__,
        "text": getattr(value, "text", None),
        "ref": getattr(value, "ref", None),
    }


def parse(content: bytes, command: CostImportDraft) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if command.sheet not in workbook.sheetnames:
            raise ProductWorkError("原工作表不存在")
        sheet = workbook[command.sheet]
        if sheet.max_row and sheet.max_row > command.header_row + 500:
            raise ProductWorkError("当前支持每批500行，请按明确范围取得分批导出原件")
        positions = {
            field: column_index_from_string(col) - 1 for field, col in command.columns.items()
        }
        lines, empty = [], []
        for number, cells in enumerate(
            sheet.iter_rows(
                min_row=command.header_row + 1,
                max_row=command.header_row + 501,
                max_col=max(positions.values()) + 1,
            ),
            command.header_row + 1,
        ):
            raw = {
                field: {
                    "cell": f"{command.columns[field]}{number}",
                    "value": cell_value(cells[i].value),
                    "type": cells[i].data_type,
                    "number_format": cells[i].number_format,
                }
                for field, i in positions.items()
            }
            if all(v["value"] in (None, "") for v in raw.values()):
                if sheet.max_row and number <= sheet.max_row:
                    empty.append(number)
                continue
            if number > command.header_row + 500:
                raise ProductWorkError("数据超过500行，未截断导入")
            lines.append({"row": number, "raw_cells": raw})
        if not lines:
            raise ProductWorkError("所选范围没有成本数据")
        return {"lines": lines, "empty_rows": empty}
    finally:
        workbook.close()


def prepare(
    c: Connection,
    context: AuthContext,
    project: UUID,
    command: CostImportDraft,
    meta: dict[str, Any],
    parsed: dict[str, Any],
) -> dict[str, Any]:
    BuildMaterialService._authorize(c, context, project, write=True)
    if metadata(document_row(c, context.tenant.id, project, command.document_id)) != meta:
        raise ProductWorkError("原件已变化，请重新读取")
    assignments = {a.row: a for a in command.assignments}
    if not assignments.keys() <= {v["row"] for v in parsed["lines"]}:
        raise ProductWorkError("人工对应必须指向本次非空原行")
    reference = f"{command.source_system}/{command.source_version}"
    lines = []
    seen = set()
    for source in parsed["lines"]:
        number = source["row"]
        raw = source["raw_cells"]
        assignment = assignments.get(number)
        value = {**source, "assignment": assignment.model_dump(mode="json") if assignment else None}
        if assignment and assignment.exclude_reason:
            lines.append({**value, "excluded": True, "error": None})
            continue
        try:
            if any(v["type"] in {"f", "e"} for v in raw.values()):
                raise ValueError("映射列有公式或错误值，请采用保留出处的ERP数值导出")
            fields = {**command.defaults, **{k: v["value"] for k, v in raw.items()}}
            for k in ("amount", "quantity", "tax_rate"):
                if isinstance(fields.get(k), bool):
                    raise ValueError("金额、数量或税率不能使用是非值")
            source_line = fields["source_line"]
            if (
                not isinstance(source_line, (str, int, float))
                or isinstance(source_line, bool)
                or source_line in (None, "")
            ):
                raise ValueError("缺少ERP唯一行标识")
            source_line = str(source_line).strip()
            if source_line in seen:
                raise ValueError("原件ERP行标识重复，不能重复计价")
            seen.add(source_line)
            for key, alias in ALIASES.items():
                fields[key] = alias.get(str(fields.get(key)), fields.get(key))
            if fields.get("tax_rate") in (None, ""):
                fields["tax_rate"] = None
            draft = PriceDraft(
                stage="actual",
                title=fields["title"],
                subsystem=fields["subsystem"],
                category=fields["category"],
                basis=fields["basis"],
                currency=fields["currency"],
                amount=fields["amount"],
                price_quantity=fields["quantity"],
                actual_quantity=fields["quantity"],
                actual_source="erp_snapshot",
                unit=fields["unit"],
                tax_basis=fields["tax_basis"],
                tax_rate=fields["tax_rate"],
                freight_basis=fields["freight_basis"],
                scope_basis=command.coverage,
                bom_item_id=assignment.bom_item_id if assignment else None,
                supplier_id=assignment.supplier_id if assignment else None,
                source_reference=reference,
                source_line=source_line,
                evidence_id=command.document_id,
                location=f"{command.sheet} 第{number}行；"
                + "；".join(f"{FIELDS[k]} {v['cell']}" for k, v in raw.items()),
                as_of=command.as_of,
                reason=command.reason,
            )
            record = uuid5(
                project,
                digest(
                    {
                        "system": command.source_system,
                        "version": command.source_version,
                        "line": source_line,
                    }
                ),
            )
            observed = input_preview(c, context, project, record, draft)
            if observed["basis"]["previous_version"] != 0:
                raise ValueError("此ERP行已有台账版本，请在原台账修订，不能批量覆盖")
            lines.append(
                {
                    **value,
                    "excluded": False,
                    "error": None,
                    "record_id": str(record),
                    "draft": draft.model_dump(mode="json"),
                    "preview": observed,
                }
            )
        except (ValueError, LookupError, PermissionError) as error:
            lines.append({**value, "excluded": False, "error": str(error)})
    value = {
        "algorithm": ALGORITHM,
        "command": command.model_dump(mode="json"),
        "document": meta,
        "lines": lines,
        "empty_rows": parsed["empty_rows"],
        "author": {"id": str(context.user.id), "name": context.user.display_name},
        "ready": any(not v["excluded"] for v in lines) and not any(v["error"] for v in lines),
    }
    return {**value, "digest": digest(value)}


class CostImportService:
    def __init__(self, engine: Engine, store: DocumentObjectStore):
        self.engine, self.store = engine, store

    def inspect(self, context: AuthContext, project: UUID, identity: UUID) -> dict[str, Any]:
        return SupplyImportService(self.engine, self.store).inspect(context, project, identity)

    def preview(
        self, context: AuthContext, project: UUID, command: CostImportDraft
    ) -> dict[str, Any]:
        meta, content = _read(self.engine, self.store, context, project, command.document_id)
        parsed = parse(content, command)
        with (
            self.engine.connect().execution_options(isolation_level="REPEATABLE READ") as c,
            c.begin(),
        ):
            return prepare(c, context, project, command, meta, parsed)

    def confirm(
        self, context: AuthContext, project: UUID, command: CostImportConfirm
    ) -> dict[str, Any]:
        signature = digest(command.model_dump(mode="json"))
        t = tables.cost_import_batches
        with self.engine.connect() as c:
            BuildMaterialService._authorize(c, context, project, write=True)
            old = (
                c.execute(
                    select(t).where(
                        t.c.tenant_id == context.tenant.id,
                        t.c.created_by_user_id == context.user.id,
                        t.c.request_id == command.request_id,
                    )
                )
                .mappings()
                .one_or_none()
            )
            if old:
                if old["project_id"] != project or old["request_digest"] != signature:
                    raise ProductWorkError("请求号已用于不同导入")
                return _plain(dict(old))
        meta, content = _read(self.engine, self.store, context, project, command.draft.document_id)
        parsed = parse(content, command.draft)
        with self.engine.begin() as c:
            NpiExecutionService._project(
                c, tenant_id=context.tenant.id, project_id=project, lock=True
            )
            BuildMaterialService._authorize(c, context, project, write=True)
            old = (
                c.execute(
                    select(t).where(
                        t.c.tenant_id == context.tenant.id,
                        t.c.created_by_user_id == context.user.id,
                        t.c.request_id == command.request_id,
                    )
                )
                .mappings()
                .one_or_none()
            )
            if old:
                if old["project_id"] != project or old["request_digest"] != signature:
                    raise ProductWorkError("请求号已用于不同导入")
                return _plain(dict(old))
            if c.execute(
                select(t.c.id).where(
                    t.c.tenant_id == context.tenant.id,
                    t.c.project_id == project,
                    t.c.source_system == command.draft.source_system,
                    t.c.source_version == command.draft.source_version,
                )
            ).first():
                raise ProductWorkError(
                    "此ERP来源版本已有导入，请查看原回执；修订原台账或使用明确的新导出版本"
                )
            if not all(
                (command.evidence_checked, command.mapping_checked, command.exclusions_checked)
            ):
                raise ProductWorkError("请核对原件权限/范围、金额数量与业务对应、排除行")
            observed = prepare(c, context, project, command.draft, meta, parsed)
            if not observed["ready"] or observed["digest"] != command.expected_digest:
                raise ProductWorkError("原行尚有错误或来源/对应已变化，请重新预览")
            batch_id = uuid4()
            receipts = []
            for line in observed["lines"]:
                if line["excluded"]:
                    continue
                saved = CostInputService(self.engine).apply_in_connection(
                    c,
                    context,
                    project,
                    UUID(line["record_id"]),
                    InputApply(
                        draft=line["draft"],
                        expected_version=0,
                        expected_digest=line["preview"]["digest"],
                        request_id=uuid5(command.request_id, f"row:{line['row']}"),
                    ),
                )
                receipts.append(
                    {
                        "row": line["row"],
                        "input_id": saved["id"],
                        "record_id": saved["record_id"],
                        "source_line": saved["payload"]["source_line"],
                        "version": saved["version"],
                    }
                )
            saved_batch = (
                c.execute(
                    t.insert()
                    .values(
                        id=batch_id,
                        tenant_id=context.tenant.id,
                        project_id=project,
                        document_id=command.draft.document_id,
                        source_system=command.draft.source_system,
                        source_version=command.draft.source_version,
                        algorithm=ALGORITHM,
                        payload=command.model_dump(mode="json"),
                        snapshot=observed,
                        receipts=receipts,
                        created_by_user_id=context.user.id,
                        request_id=command.request_id,
                        request_digest=signature,
                        created_at=func.clock_timestamp(),
                    )
                    .returning(t)
                )
                .mappings()
                .one()
            )
            for receipt in receipts:
                c.execute(
                    tables.cost_import_rows.insert().values(
                        id=uuid4(),
                        tenant_id=context.tenant.id,
                        project_id=project,
                        batch_id=batch_id,
                        row_number=receipt["row"],
                        input_id=UUID(receipt["input_id"]),
                    )
                )
            NpiExecutionService._record_change(
                c,
                tenant_id=context.tenant.id,
                project_id=project,
                actor_user_id=context.user.id,
                record_id=batch_id,
                record_type="cost_import",
                action="captured",
                before=None,
                after=_plain(dict(saved_batch)),
            )
            return _plain(dict(saved_batch))

    def get(
        self,
        context: AuthContext,
        project: UUID,
        identity: UUID | None = None,
        before: UUID | None = None,
        search: str | None = None,
    ) -> dict[str, Any]:
        with (
            self.engine.connect().execution_options(isolation_level="REPEATABLE READ") as c,
            c.begin(),
        ):
            BuildMaterialService._authorize(c, context, project)
            t = tables.cost_import_batches
            d = tables.documents
            columns = (
                list(t.c)
                if identity
                else [
                    t.c.id,
                    t.c.source_system,
                    t.c.source_version,
                    t.c.document_id,
                    t.c.created_at,
                ]
            )
            q = select(*columns).where(
                t.c.tenant_id == context.tenant.id, t.c.project_id == project
            )
            if identity:
                q = q.where(t.c.id == identity)
            if before:
                cursor = row(c, context.tenant.id, project, t, before)
                q = q.where(
                    (t.c.created_at < cursor["created_at"])
                    | ((t.c.created_at == cursor["created_at"]) & (t.c.id < before))
                )
            records = list(
                c.execute(q.order_by(t.c.created_at.desc(), t.c.id.desc()).limit(21)).mappings()
            )
            values = []
            for record in records[:20]:
                value = _plain(dict(record))
                if identity:
                    try:
                        value["document_current"] = (
                            metadata(
                                document_row(c, context.tenant.id, project, record["document_id"])
                            )
                            == record["snapshot"]["document"]
                        )
                    except (ValueError, LookupError, PermissionError):
                        value["document_current"] = False
                    ids = [UUID(r["record_id"]) for r in record["receipts"]]
                    current = (
                        c.execute(
                            current_query(context.tenant.id, project).where(
                                tables.cost_inputs.c.record_id.in_(ids)
                            )
                        )
                        .mappings()
                        .all()
                    )
                    value["current_inputs"] = [
                        {
                            "record_id": str(v["record_id"]),
                            "id": str(v["id"]),
                            "version": v["version"],
                            "status": v["status"],
                        }
                        for v in current
                    ]
                values.append(value)
            dq = select(d.c.id, d.c.filename).where(
                d.c.tenant_id == context.tenant.id,
                d.c.project_id == project,
                d.c.media_type == MEDIA,
            )
            if search:
                needle = (
                    "%"
                    + search.strip().replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
                    + "%"
                )
                dq = dq.where(d.c.filename.ilike(needle, escape="\\"))
            documents = list(
                c.execute(dq.order_by(d.c.created_at.desc(), d.c.id).limit(201)).mappings()
            )
            return {
                "items": values,
                "next_cursor": str(records[19]["id"]) if len(records) > 20 else None,
                "documents": [_plain(dict(v)) for v in documents[:200]],
                "documents_more": len(documents) > 200,
                "fields": FIELDS,
            }
