"""Actual internal definition bytes and frozen download receipts; synthetic business data."""

# ruff: noqa: F811
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.exc import DBAPIError

from app.db.schema import tables
from app.domain.evidence import EvidenceService
from app.exports.definition_documents import MIMES, Format
from tests.test_product_manager_loop_api import command
from tests.test_product_planning_findings_api import post, release_definition, setup_candidate
from tests.test_product_work_preparation_api import space
from tests.test_test_runs_api import Scope
from tests.test_test_runs_api import client as client
from tests.test_test_runs_api import isolated_test_run_database as isolated_test_run_database


def extracted_text(content: bytes, format: Format) -> str:
    source = BytesIO(content)
    if format == "docx":
        return "\n".join(p.text for p in Document(source).paragraphs)
    if format == "xlsx":
        workbook = load_workbook(source, data_only=False)
        try:
            return "\n".join(str(cell.value) for sheet in workbook for row in sheet for cell in row)
        finally:
            workbook.close()
    if format == "pptx":
        return "\n".join(
            shape.text
            for slide in Presentation(source).slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
    return "\n".join(page.extract_text() for page in PdfReader(source).pages)


def test_published_definition_exports_keep_bytes_and_recheck_current_release(
    client: TestClient,
    isolated_test_run_database: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    scope = Scope(client, isolated_test_run_database)
    try:
        base, headers, definition, calls = setup_candidate(
            client,
            scope,
            tmp_path,
            monkeypatch,
            "product_definition",
            revision=True,
            seed_only=True,
        )
        assert calls == []  # The fixture never invokes a model or installs a qualification.
        options = post(client, base, headers, "/definition-export-options", {"search": "合成"})
        assert definition["id"] in {item["id"] for item in options["items"]}
        receipts: list[dict[str, Any]] = []
        originals: dict[str, bytes] = {}
        saved_command: dict[str, Any] = {}
        formats: tuple[Format, ...] = ("docx", "xlsx", "pptx", "pdf")
        for format in formats:
            selected = dict(definition_id=definition["id"], format=format, reason="合成内部核对")
            preview = post(client, base, headers, "/definition-export-preview", selected)
            assert preview["source"]["status_label"] == "当前已发布定义"
            assert len(preview["source"]["reviews"]) == 5
            assert preview["source"]["definition"]["strategy_ref"] == definition["strategy_ref"]
            assert (
                preview["source"]["lineage"]["strategy"]["reference"] == definition["strategy_ref"]
            )
            assert preview["source"]["template_version"] == "internal-definition.v1"
            assert preview["source"]["source_policy"] == "definition-export.v2"
            saved_command = {
                **selected,
                "request_id": str(uuid4()),
                "expected_digest": preview["digest"],
            }
            receipt = post(client, base, headers, "/definition-export-save", saved_command)
            assert post(client, base, headers, "/definition-export-save", saved_command) == receipt
            download = client.get(
                base + f"/definition-exports/{receipt['id']}/download", headers=headers
            )
            assert download.status_code == 200, download.text
            assert download.headers["content-type"] == MIMES[format]
            assert download.headers["x-content-type-options"] == "nosniff"
            assert "no-store" in download.headers["cache-control"]
            assert receipt["filename"] in download.headers["content-disposition"]
            assert len(download.content) == receipt["byte_size"]
            assert sha256(download.content).hexdigest() == receipt["content_digest"]
            text = "".join(extracted_text(download.content, format).split())
            for required in (
                definition["title"],
                definition["id"],
                receipt["id"],
                "当前已发布定义",
                "pressure",
                "kPa",
                "25℃",
                "合成判据",
                "专业协审历史",
                "approved",
                definition["strategy_ref"]["ref_id"],
            ):
                assert required in text, (format, required)
            # Presentation layout wraps long hashes; the underlying source stays exact.
            assert definition["strategy_ref"]["digest"] in "".join(text.split())
            originals[receipt["id"]] = download.content
            receipts.append(receipt)

        with scope.engine.connect() as c:
            assert (
                c.execute(select(func.count()).select_from(tables.definition_exports)).scalar_one()
                == 4
            )
        changed = client.post(
            base + "/definition-export-save",
            headers=headers,
            json={**saved_command, "reason": "同一请求篡改用途"},
        )
        assert changed.status_code == 409, changed.text
        history = post(
            client, base, headers, "/definition-export-history", {"definition_id": definition["id"]}
        )
        assert len(history["items"]) == 4 and all(
            item["sources_current"] for item in history["items"]
        )
        other_base = f"/api/product-work/spaces/{space(client, headers)}"
        denied = client.get(
            other_base + f"/definition-exports/{receipts[0]['id']}/download", headers=headers
        )
        assert denied.status_code == 409 and receipts[0]["filename"] not in denied.text

        post(
            client,
            base,
            headers,
            f"/definitions/{definition['id']}/actions",
            {
                **command(definition),
                "action": "withdraw",
                "reason": "合成撤回验证",
                "expected_release_id": definition["current_release"]["id"],
            },
        )
        stale = client.post(
            base + "/definition-export-save",
            headers=headers,
            json={**saved_command, "request_id": str(uuid4())},
        )
        assert stale.status_code == 409 and "重新预览" in stale.text
        history = post(
            client, base, headers, "/definition-export-history", {"definition_id": definition["id"]}
        )
        assert len(history["items"]) == 4 and not any(
            item["sources_current"] for item in history["items"]
        )
        for receipt in receipts:
            historical = client.get(
                base + f"/definition-exports/{receipt['id']}/download", headers=headers
            )
            assert historical.status_code == 200 and historical.content == originals[receipt["id"]]
        preview = post(client, base, headers, "/definition-export-preview", selected)
        assert preview["source"]["status_label"] == "历史已发布定义（当前不生效）"
        with pytest.raises(DBAPIError), scope.engine.begin() as c:
            c.execute(tables.definition_exports.update().values(reason="不可覆盖历史"))
        with scope.engine.connect() as c:
            assert (
                c.execute(select(func.count()).select_from(tables.model_invocations)).scalar_one()
                == 0
            )
    finally:
        scope.dispose()


def test_derived_export_keeps_locked_content_after_parent_superseded_and_withdrawn(
    client: TestClient,
    isolated_test_run_database: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    scope = Scope(client, isolated_test_run_database)
    try:
        base, headers, parent, calls = setup_candidate(
            client,
            scope,
            tmp_path,
            monkeypatch,
            "product_definition",
            revision=True,
            seed_only=True,
        )
        assert calls == []
        child = post(
            client,
            base,
            headers,
            f"/definitions/{parent['id']}/derive",
            {
                **command(parent),
                "code": "EXPORT-DERIVED",
                "name": "合成派生导出产品",
                "reason": "合成锁定来源核对",
            },
            201,
        )
        content = deepcopy(child["content"])
        content["reviewers"] = {d: str(scope.owner_id) for d in ["id", "md", "ee", "fw", "dqe"]}
        content["items"][0].update(state="confirmed", owner_user_id=str(scope.owner_id))
        child = post(
            client,
            base,
            headers,
            f"/definitions/{child['id']}/changes",
            {**command(child), "content": content, "reason": "合成具名确认，非真实协审"},
        )
        child = release_definition(client, base, headers, child)
        locked = deepcopy(child["content"])
        selected = {"definition_id": child["id"], "format": "docx", "reason": "合成内部核对"}
        original = post(client, base, headers, "/definition-export-preview", selected)
        assert original["source"]["lineage"]["parent_definition"]["status"] == "current"
        saved_command = {
            **selected,
            "request_id": str(uuid4()),
            "expected_digest": original["digest"],
        }
        saved = post(client, base, headers, "/definition-export-save", saved_command)
        url = base + f"/definition-exports/{saved['id']}/download"
        frozen = client.get(url, headers=headers)
        assert frozen.status_code == 200

        content = deepcopy(parent["content"])
        content["items"][0]["target_value"] = "110"
        content["change_reason"] = "合成父版替代"
        incoming = post(
            client,
            base,
            headers,
            "/definitions",
            {
                "request_id": str(uuid4()),
                "product_id": parent["product_id"],
                "strategy_id": parent["strategy_id"],
                "supersedes_id": parent["id"],
                "content": content,
            },
            201,
        )
        incoming = release_definition(client, base, headers, incoming)
        for expected in ["superseded", "withdrawn"]:
            if expected == "withdrawn":
                post(
                    client,
                    base,
                    headers,
                    f"/definitions/{incoming['id']}/actions",
                    {
                        **command(incoming),
                        "action": "withdraw",
                        "reason": "合成父产品撤回核对",
                        "expected_release_id": incoming["current_release"]["id"],
                    },
                )
            current = post(client, base, headers, "/definition-export-preview", selected)
            source = current["source"]
            assert source["status_label"] == "当前已发布定义"
            assert source["definition"]["content"] == locked
            assert source["definition"]["base_definition_ref"]["ref_id"] == parent["id"]
            assert source["lineage"]["parent_definition"]["status"] == expected
            assert any("父" in warning for warning in source["warnings"])
            assert current["digest"] != original["digest"]
            post(
                client,
                base,
                headers,
                "/definition-export-save",
                {**saved_command, "request_id": str(uuid4())},
                409,
            )
            assert post(client, base, headers, "/definition-export-save", saved_command) == saved
            historical = client.get(url, headers=headers)
            assert historical.status_code == 200 and historical.content == frozen.content
        with scope.engine.connect() as c:
            assert (
                c.execute(select(func.count()).select_from(tables.model_invocations)).scalar_one()
                == 0
            )
    finally:
        scope.dispose()


def test_export_download_rechecks_original_evidence_after_source_is_disabled(
    client: TestClient,
    isolated_test_run_database: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    scope = Scope(client, isolated_test_run_database)
    try:
        base, headers, original, _ = setup_candidate(
            client,
            scope,
            tmp_path,
            monkeypatch,
            "product_definition",
            revision=True,
            seed_only=True,
        )
        evidence = EvidenceService(scope.engine)
        source = evidence.create_source(
            tenant_id=scope.tenant_id,
            actor_user_id=scope.owner_id,
            name="合成导出依据",
            source_type="file",
            allowed_use="internal_analysis",
            collection_frequency="on_demand",
            retention_policy="standard",
            trust_tier="unknown",
        )
        snapshot = evidence.capture(
            tenant_id=scope.tenant_id,
            actor_user_id=scope.owner_id,
            source_id=source.id,
            content="合成原文，不应自动附入规格导出文件。",
        )
        content = deepcopy(original["content"])
        content["change_reason"] = "补充合成来源"
        content["target_cost"]["sources"] = [
            {
                "tenant_id": str(scope.tenant_id),
                "kind": "evidence_snapshot",
                "ref_id": str(snapshot.id),
                "as_of": snapshot.captured_at.isoformat(),
                "digest": snapshot.content_digest,
            }
        ]
        draft = post(
            client,
            base,
            headers,
            "/definitions",
            {
                "request_id": str(uuid4()),
                "product_id": original["product_id"],
                "strategy_id": original["strategy_id"],
                "supersedes_id": original["id"],
                "content": content,
            },
            201,
        )
        selected = dict(definition_id=draft["id"], format="docx", reason="内部草稿核对")
        preview = post(client, base, headers, "/definition-export-preview", selected)
        assert preview["source"]["status_label"] == "草稿 · 尚未发布"
        request = {**selected, "request_id": str(uuid4()), "expected_digest": preview["digest"]}
        saved = post(client, base, headers, "/definition-export-save", request)
        url = base + f"/definition-exports/{saved['id']}/download"
        download = client.get(url, headers=headers)
        assert download.status_code == 200
        text = extracted_text(download.content, "docx")
        assert "草稿 · 尚未发布" in text and str(snapshot.id) in text
        assert "合成原文，不应自动附入规格导出文件。" not in text
        with scope.engine.begin() as c:
            c.execute(
                tables.research_sources.update()
                .where(tables.research_sources.c.id == source.id)
                .values(status="disabled")
            )
        for path, body in [
            ("/definition-export-preview", selected),
            ("/definition-export-save", request),
        ]:
            rejected = client.post(base + path, headers=headers, json=body)
            assert rejected.status_code == 409, rejected.text
        denied = client.get(url, headers=headers)
        assert denied.status_code == 409 and "content-disposition" not in denied.headers
        history = post(
            client, base, headers, "/definition-export-history", {"definition_id": draft["id"]}
        )
        assert len(history["items"]) == 1 and not history["items"][0]["sources_current"]
        with scope.engine.connect() as c:
            assert (
                c.execute(select(func.count()).select_from(tables.definition_exports)).scalar_one()
                == 1
            )
    finally:
        scope.dispose()
