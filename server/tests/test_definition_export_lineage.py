"""Synthetic unit checks of real projection/rendering; no database or API fixtures."""

from collections.abc import Iterator
from contextlib import contextmanager
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace
from typing import Any, Literal, cast
from uuid import UUID

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
from sqlalchemy import Connection as SQLConnection
from sqlalchemy import Engine as SQLEngine
from sqlalchemy.sql import Executable

from app.auth.repository import AuthContext
from app.contracts.common_objects import SourceRef
from app.domain import definition_exports as exports
from app.domain.product_workspace import ProductWorkAccessError, digest
from app.exports.definition_documents import Format, render, sections

TENANT, SPACE, DEFINITION, PRODUCT, STRATEGY, PARENT, OTHER = [UUID(int=i) for i in range(1, 8)]
DIGEST = "a" * 64


class Result:
    def __init__(self, value: Any) -> None:
        self.value = value

    def mappings(self) -> "Result":
        return self

    def all(self) -> Any:
        return self.value

    def one(self) -> Any:
        return self.value

    def one_or_none(self) -> Any:
        return self.value


class Connection:
    def __init__(self, values: list[Any]) -> None:
        self.values = list(values)

    def execute(self, statement: Executable) -> Result:
        assert str(statement), "Real SQL boundary must receive a statement"
        assert self.values, "Unexpected SQL call"
        return Result(self.values.pop(0))


def reference(kind: str, identity: UUID) -> dict[str, Any]:
    return dict(
        tenant_id=str(TENANT),
        kind=kind,
        ref_id=str(identity),
        version=2,
        digest=DIGEST,
        as_of="2026-09-30T00:00:00Z",
    )


def setup(
    monkeypatch: pytest.MonkeyPatch,
    *,
    parent: bool = True,
    release_action: Literal["publish", "withdraw"] | None = "publish",
    parent_latest: UUID = PARENT,
    parent_version: int = 2,
    parent_digest: str = DIGEST,
    strategy_latest: UUID = STRATEGY,
    strategy_version: int = 2,
    strategy_digest: str = DIGEST,
    deny: Literal["parent", "strategy"] | None = None,
) -> tuple[Connection, AuthContext, dict[str, Any]]:
    fact = dict(text="Synthetic assumption", fact_kind="hypothesis", sources=[])
    content: dict[str, Any] = {
        key: []
        for key in (
            "target_users",
            "scenarios",
            "included_features",
            "excluded_features",
            "region_constraints",
            "service_requirements",
            "success_criteria",
            "items",
        )
    }
    content.update(
        target_cost=deepcopy(fact),
        price_position=deepcopy(fact),
        launch_window=deepcopy(fact),
        change_reason="Synthetic unit only",
    )
    row: dict[str, Any] = dict(
        id=DEFINITION,
        title="Synthetic lineage",
        product_id=PRODUCT,
        strategy_id=STRATEGY,
        strategy_ref=reference("product_strategy", STRATEGY),
        series_id=OTHER,
        revision_number=1,
        version=2,
        state="published",
        review_round=1,
        review_digest=DIGEST,
        approval=None,
        published_at=None,
        content=content,
        base_definition_id=PARENT if parent else None,
        base_definition_ref=reference("product_definition", PARENT),
        parent_updates=[],
    )
    parent_row = dict(product_id=OTHER, version=parent_version, review_digest=parent_digest)
    strategy = dict(
        id=STRATEGY,
        title="Locked strategy",
        series_id=OTHER,
        product_line_id=PRODUCT,
        revision_number=1,
        version=strategy_version,
        state="published",
        approval={"payload_digest": strategy_digest},
    )

    def read(c: SQLConnection, context: AuthContext, space: UUID, identity: UUID) -> dict[str, Any]:
        if identity == PARENT:
            if deny == "parent":
                raise ProductWorkAccessError("parent scope denied")
            return parent_row
        assert identity == DEFINITION
        return row

    def line(*args: Any) -> None:
        if deny == "strategy":
            raise ProductWorkAccessError("strategy scope denied")

    def release(c: SQLConnection, context: AuthContext, product: UUID) -> dict[str, Any] | None:
        if product == PRODUCT:
            return dict(id=OTHER, definition_id=DEFINITION, action="publish")
        return (
            None
            if release_action is None
            else dict(definition_id=parent_latest, action=release_action)
        )

    def sources(
        c: SQLConnection, context: AuthContext, refs: tuple[SourceRef, ...]
    ) -> list[dict[str, Any]]:
        assert refs == (), "Hypotheses require no research evidence"
        return []

    monkeypatch.setattr(exports, "_space", lambda *a, **kw: None)
    monkeypatch.setattr(exports, "_row", read)
    monkeypatch.setattr(exports, "_line", line)
    monkeypatch.setattr(exports, "_release", release)
    monkeypatch.setattr(exports, "_sources", sources)
    monkeypatch.setattr(exports, "evidence_quality", lambda *a: [])
    c = Connection(
        [
            [],
            [],
            dict(id=PRODUCT, name="Synthetic product", code="SYN"),
            strategy,
            dict(id=strategy_latest, revision_number=2),
        ]
    )
    return (
        c,
        cast(
            AuthContext,
            SimpleNamespace(tenant=SimpleNamespace(id=TENANT), user=SimpleNamespace(id=OTHER)),
        ),
        row,
    )


def extract(data: bytes, format: Format) -> str:
    stream = BytesIO(data)
    if format == "docx":
        return "\n".join(p.text for p in Document(stream).paragraphs)
    if format == "xlsx":
        sheet = load_workbook(stream).active
        assert sheet is not None
        return "\n".join(str(cell.value) for row in sheet for cell in row)
    if format == "pptx":
        return "".join(
            slide.shapes[1].text.replace("\n", "")
            for slide in Presentation(stream).slides
            if len(slide.shapes) > 1
        )
    return "\n".join(page.extract_text() for page in PdfReader(stream).pages)


@pytest.mark.parametrize("format", ["docx", "pdf", "xlsx", "pptx"])
def test_hypothesis_projection_preserves_locked_strategy_in_all_bytes(
    monkeypatch: pytest.MonkeyPatch, format: Format
) -> None:
    c, context, original = setup(monkeypatch)
    snapshot = exports.source(cast(SQLConnection, c), context, SPACE, DEFINITION)
    assert not c.values
    assert snapshot["sources"] == []
    assert snapshot["definition"]["strategy_ref"] == original["strategy_ref"]
    assert snapshot["definition"]["content"] == original["content"]
    assert snapshot["digest"] == digest({k: v for k, v in snapshot.items() if k != "digest"})
    changed = deepcopy(snapshot)
    changed["lineage"]["strategy"]["status"] = "superseded"
    assert digest({k: v for k, v in changed.items() if k != "digest"}) != snapshot["digest"]
    text = extract(render(snapshot, format, {"编号": "synthetic-lineage"}), format)
    assert str(STRATEGY) in text and DIGEST in text and "product_strategy" in text
    assert "Synthetic assumption" in text and "current" in text


@pytest.mark.parametrize(
    "options,status",
    [
        ({"release_action": "withdraw"}, "withdrawn"),
        ({"parent_latest": OTHER}, "superseded"),
        ({"release_action": None}, "no_current_release"),
        ({"parent_version": 3}, "unavailable"),
        ({"parent_digest": "b" * 64}, "unavailable"),
    ],
)
def test_parent_status_changes_do_not_replace_locked_content(
    monkeypatch: pytest.MonkeyPatch, options: dict[str, Any], status: str
) -> None:
    c, context, original = setup(monkeypatch, **options)
    snapshot = exports.source(cast(SQLConnection, c), context, SPACE, DEFINITION)
    assert snapshot["lineage"]["parent_definition"]["status"] == status
    assert snapshot["warnings"]
    assert snapshot["definition"]["content"] == original["content"]
    assert snapshot["definition"]["base_definition_ref"] == original["base_definition_ref"]
    assert snapshot["status_label"] == "当前已发布定义"


@pytest.mark.parametrize(
    "options,status",
    [
        ({"strategy_latest": OTHER}, "superseded"),
        ({"strategy_version": 3}, "unavailable"),
        ({"strategy_digest": "b" * 64}, "unavailable"),
    ],
)
def test_new_or_mismatched_strategy_never_substitutes_fixed_reference(
    monkeypatch: pytest.MonkeyPatch, options: dict[str, Any], status: str
) -> None:
    c, context, original = setup(monkeypatch, **options)
    snapshot = exports.source(cast(SQLConnection, c), context, SPACE, DEFINITION)
    assert snapshot["lineage"]["strategy"]["status"] == status
    assert snapshot["warnings"]
    assert snapshot["definition"]["strategy_ref"] == original["strategy_ref"]


@pytest.mark.parametrize("deny", ["parent", "strategy"])
def test_denied_business_source_scope_prevents_projection_and_render(
    monkeypatch: pytest.MonkeyPatch, deny: Literal["parent", "strategy"]
) -> None:
    c, context, _ = setup(monkeypatch, deny=deny)
    with pytest.raises(ProductWorkAccessError):
        render(exports.source(cast(SQLConnection, c), context, SPACE, DEFINITION), "docx", {})


def test_legacy_projection_still_renders_without_lineage(monkeypatch: pytest.MonkeyPatch) -> None:
    c, context, _ = setup(monkeypatch, parent=False)
    snapshot = exports.source(cast(SQLConnection, c), context, SPACE, DEFINITION)
    snapshot.pop("lineage")
    snapshot["definition"].pop("strategy_ref")
    snapshot["template_version"] = "internal-definition.v1"
    assert "internal-definition.v1" in "\n".join(body for _, body in sections(snapshot))
    formats: tuple[Format, ...] = ("docx", "pdf", "xlsx", "pptx")
    for format in formats:
        assert "Synthetic assumption" in extract(
            render(snapshot, format, {"编号": "synthetic-legacy"}), format
        )


class Engine:
    """Only transaction/connection boundaries; service SQL still uses the queued result double."""

    def __init__(self, connection: Connection) -> None:
        self.connection = connection

    @contextmanager
    def connect(self) -> Iterator[SQLConnection]:
        yield cast(SQLConnection, self.connection)

    @contextmanager
    def begin(self) -> Iterator[SQLConnection]:
        yield cast(SQLConnection, self.connection)


def stored_export(
    monkeypatch: pytest.MonkeyPatch,
) -> tuple[dict[str, Any], exports.ExportSave]:
    c, context, _ = setup(monkeypatch)
    snapshot = exports.source(cast(SQLConnection, c), context, SPACE, DEFINITION)
    command = exports.ExportSave(
        definition_id=DEFINITION,
        format="docx",
        reason="Synthetic saved export",
        request_id=OTHER,
        expected_digest=DIGEST,
    )
    content = render(snapshot, "docx", {"编号": "synthetic-retained-export"})
    saved = dict(
        id=OTHER,
        definition_id=DEFINITION,
        format="docx",
        template_version=snapshot["template_version"],
        filename="synthetic-retained.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        byte_size=len(content),
        content_digest=sha256(content).hexdigest(),
        content=content,
        snapshot=deepcopy(snapshot),
        source_digest=snapshot["digest"],
        reason=command.reason,
        actor_user_id=context.user.id,
        created_at="2026-09-30T00:00:00Z",
        request_digest=digest({"space": str(SPACE), "command": command.model_dump(mode="json")}),
    )
    return saved, command


@pytest.mark.parametrize("operation", ["download", "save_retry"])
@pytest.mark.parametrize("deny", ["parent", "strategy"])
def test_retained_export_reauthorizes_business_sources_before_return(
    monkeypatch: pytest.MonkeyPatch,
    operation: Literal["download", "save_retry"],
    deny: Literal["parent", "strategy"],
) -> None:
    saved, command = stored_export(monkeypatch)
    before = deepcopy(saved)
    c, context, child = setup(monkeypatch, deny=deny)
    # The child remains readable; rejection must come from its locked business source.
    read_child = cast(Any, exports)._row
    assert read_child(cast(SQLConnection, c), context, SPACE, DEFINITION) is child
    c.values = [saved, *c.values[3:]]
    service = exports.DefinitionExportService(cast(SQLEngine, Engine(c)))
    with pytest.raises(ProductWorkAccessError, match=f"{deny} scope denied"):
        if operation == "download":
            service.download(context, SPACE, OTHER)
        else:
            service.save(context, SPACE, command)
    assert saved == before


def test_authorized_download_after_parent_changes_returns_frozen_bytes(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    saved, _ = stored_export(monkeypatch)
    before = deepcopy(saved)
    c, context, _ = setup(monkeypatch, release_action="withdraw", strategy_latest=OTHER)
    c.values = [saved, *c.values[3:]]
    service = exports.DefinitionExportService(cast(SQLEngine, Engine(c)))
    content, mime_type, filename = service.download(context, SPACE, OTHER)
    assert content == before["content"]
    assert sha256(content).hexdigest() == before["content_digest"]
    assert mime_type == before["mime_type"] and filename == before["filename"]
    assert saved == before and not c.values
    assert saved["snapshot"]["lineage"]["parent_definition"]["status"] == "current"
