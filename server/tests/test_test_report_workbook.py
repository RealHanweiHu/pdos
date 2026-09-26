"""Report layouts are evidence, not a universal product test checklist."""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

from app.domain.test_report_workbook import parse_workbook


def form_report() -> bytes:
    book = Workbook()
    cover = book.active
    assert cover is not None
    cover["A1"] = "Example company"
    cover["A2"] = "整机性能测试报告"
    directory = book.create_sheet("测试目录")
    directory.append(["序号", "测试项", "判定"])
    directory.append([1, "气压精度", "合格"])
    for title, verdict in [("气压精度", "合格"), ("工作噪音", "参考"), ("照明", "不适用")]:
        sheet = book.create_sheet(title + " ")
        sheet["A1"] = title + "测试"
        sheet["A2"] = "样机型号"
        sheet["D2"] = "BICYCLE-REF"
        sheet["G2"] = "测试项目"
        sheet["J2"] = title
        sheet["M2"] = "测试日期"
        sheet["P2"] = datetime(2025, 11, 5)
        sheet["A5"] = "测试方法"
        sheet["D5"] = "执行测试并记录"
        sheet["A6"] = "判定标准"
        sheet["D6"] = "参照规格书"
        sheet["A7"] = "测试记录"
        sheet["A8"] = "序号"
        sheet["D8"] = "实测记录"
        sheet["A9"] = 1
        sheet["D9"] = "83.5 dB"
        sheet["A10"] = "结论"
        sheet["D10"] = verdict
        sheet.merge_cells("D9:F9")
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_forms_scan_actual_sheets_keep_sources_dates_and_non_pass_verdicts() -> None:
    report = parse_workbook(form_report())
    assert report.layout == "test_form_v1"
    assert len(report.rows) == 3  # directory deliberately omits two tests
    assert [row.values["测试结果"] for row in report.rows] == ["合格", "参考", "不适用"]
    row = report.rows[0]
    assert row.source["sheet"] == "气压精度 "
    assert row.source["result_cell"] == "D10"
    assert row.source["reported_date"] == "2025-11-05"
    assert row.source["model"] == "BICYCLE-REF"
    assert row.values["样机SN"] == ""  # a model or a row number is not an SN
    assert row.source["records"][1]["cells"]["D9"] == "83.5 dB"


def test_missing_date_and_blank_observations_are_preserved_as_missing() -> None:
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(form_report()))
    sheet = book["气压精度 "]
    sheet["P2"] = None
    sheet["D9"] = None
    buffer = BytesIO()
    book.save(buffer)
    row = parse_workbook(buffer.getvalue()).rows[0]
    assert row.source["reported_date"] is None
    assert row.source["has_observations"] is False
    assert row.values["测试结果"] == "合格"


def test_flat_workbooks_keep_sheet_and_physical_row_numbers() -> None:
    book = Workbook()
    first = book.active
    assert first is not None
    for sheet in [first, book.create_sheet("second")]:
        sheet.append(["用例名称", "测试结果"])
        sheet.append([None, None])
        sheet.append(["Leak", "PASS"])
    buffer = BytesIO()
    book.save(buffer)
    report = parse_workbook(buffer.getvalue())
    assert report.layout == "flat_v1"
    assert len(report.rows) == 2
    assert [row.source["row"] for row in report.rows] == [3, 3]
    assert report.rows[1].source["sheet"] == "second"
