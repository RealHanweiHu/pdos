"""Stage integration of retained report bytes across all supported formats."""

from io import BytesIO
from typing import Literal

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
from pypdf.generic import DictionaryObject

from app.exports.definition_documents import Format, render_sections


@pytest.mark.parametrize(
    "original",
    [
        "".join(f"{i:03d}：协审原记录，后半部分同样必须可读。\n" for i in range(110)),
        "中文规格与工况" * 700,
    ],
)
def test_workbook_breaks_tall_records_into_visible_continuations(original: str) -> None:
    content = render_sections(
        "合成规格",
        "草稿 · 内部核对",
        [("完整来源", original)],
        "xlsx",
        {"编号": "synthetic-visible-rows"},
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook.active
    assert sheet is not None
    source_rows = [row for row in sheet if str(row[0].value).startswith("完整来源")]
    assert len(source_rows) > 2
    retained: list[str] = []
    for row in source_rows:
        cell_value = row[1].value
        assert isinstance(cell_value, str)
        retained.append(cell_value)
    assert "".join(retained) == original
    # Explicit bounded heights are required; relying on auto-fit can paint
    # newline-heavy records outside a printed page despite retained XML text.
    for row in source_rows:
        row_number = row[0].row
        assert row_number is not None
        height = sheet.row_dimensions[row_number].height
        assert height is not None and 24 <= height <= 216
        assert all(cell.alignment.wrap_text for cell in row)
    assert sheet.print_title_rows == "$1:$1"
    assert sheet.page_setup.fitToHeight == 0


def test_workbook_long_section_label_is_retained_in_visible_body() -> None:
    heading = "长中文规格标题与工况说明" * 20
    content = render_sections(
        "合成规格", "草稿", [(heading, "固定原值")], "xlsx", {"编号": "synthetic-label"}
    )
    sheet = load_workbook(BytesIO(content)).active
    assert sheet is not None
    rows = [row for row in sheet if str(row[0].value).startswith("长中文规格标题")]
    assert "".join(str(row[1].value or "") for row in rows) == heading + "\n\n固定原值"
    for row in rows:
        row_number = row[0].row
        assert row_number is not None
        height = sheet.row_dimensions[row_number].height
        assert height is not None and height <= 216


def test_pdf_mixed_text_uses_portable_ascii_font_in_body_and_footer() -> None:
    receipt_id = "08d499bc-6060-47af-91ee-2603553e3cc1"
    digest = "a7b3c9d2" * 8
    title = '定义 PDOS <A&B> "v1"'
    timestamp = "2026-09-30T10:24:31.123456+00:00"
    symbols = "· ± ° – — ‘西文’ “标点” €"
    source = (
        f'中文字段：{{"id": "{receipt_id}", "sha256": "{digest}"}}\n时间 {timestamp}'
        f"\n工况 25 ± 2 °C · 25℃ / 32℉ · {symbols} / <>&"
    )
    pdf = PdfReader(
        BytesIO(
            render_sections(
                title, "已发布 PM", [("规格 PM-001", source * 60)], "pdf", {"编号": receipt_id}
            )
        )
    )
    fragments: list[tuple[str, str]] = []

    def collect(
        text: str,
        cm: list[float],
        tm: list[float],
        font: DictionaryObject | None,
        size: float,
    ) -> None:
        if text.strip():
            fragments.append((text, str(font.get("/BaseFont")) if font else ""))

    extracted = "".join(page.extract_text(visitor_text=collect) for page in pdf.pages)
    compact = "".join(extracted.split())
    for expected in (title, "规格 PM-001", receipt_id, digest, timestamp, symbols, "/ <>&"):
        assert "".join(expected.split()) in compact
    assert len(pdf.pages) > 1  # Exercise wrapping and every page's footer.
    assert any("中文字段" in text and font == "/STSong-Light" for text, font in fragments)
    # UUIDs, JSON keys, timestamps and heading/footer identifiers must never
    # use the viewer-dependent Latin glyph metrics of the Chinese CID font.
    assert all(
        font == "/Helvetica"
        for text, font in fragments
        if any((char.isascii() and char.isalnum()) or char in "·±°–—‘’“”€" for char in text)
    )
    for symbol in "·±°–—‘’“”€":
        assert any(symbol in text and font == "/Helvetica" for text, font in fragments)
    assert "25°C/32°F" in compact
    assert "℃" not in extracted and "℉" not in extracted
    assert any("25°C / 32°F" in text and font == "/Helvetica" for text, font in fragments)
    for page in pdf.pages:
        text = "".join(page.extract_text().split())
        assert f"已发布PM·内部·{receipt_id}·" in text


@pytest.mark.parametrize("format", ["docx", "xlsx", "pptx", "pdf"])
@pytest.mark.parametrize("audience", ["internal", "customer"])
def test_generated_report_retains_long_sources_and_receipt(
    format: Format, audience: Literal["internal", "customer"]
) -> None:
    # Synthetic text exceeds each renderer's cell/paragraph/slide chunk size.
    original = "来源边界BEGIN " + "逐项原文甲乙丙1234567890" * 420 + " 来源边界END"
    formula = '=HYPERLINK("https://example.invalid/never-fetch", "原值")'
    receipt = {"编号": "synthetic-stage-export", "资料": "仅供隔离集成验证"}
    status = "历史版本 · 非真实业务结论"
    content = render_sections(
        "合成规格与质量报告",
        status,
        [("完整来源", original), ("原值记录", formula), ("缺项", "未记录；不补零")],
        format,
        receipt,
        audience=audience,
    )
    source = BytesIO(content)
    source_text: str | None = None
    if format == "docx":
        document = Document(source)
        extracted = "\n".join(p.text for p in document.paragraphs)
        footer = document.sections[0].footer.paragraphs[0].text
        assert receipt["编号"] in footer
    elif format == "xlsx":
        workbook = load_workbook(source, data_only=False)
        sheet = workbook.active
        assert sheet is not None
        extracted = "\n".join(str(cell.value) for row in sheet for cell in row)
        formula_cells = [cell for row in sheet for cell in row if cell.value == formula]
        assert len(formula_cells) == 1
        assert formula_cells[0].data_type == "s"
        assert formula_cells[0].hyperlink is None
        source_cells = [
            str(row[1].value) for row in sheet if str(row[0].value).startswith("完整来源")
        ]
        assert "".join(source_cells) == original
        source_text = "".join(source_cells)
    elif format == "pptx":
        presentation = Presentation(source)
        assert len(presentation.slides) > 4
        extracted = "\n".join(
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        source_text = "".join(
            slide.shapes[1].text
            for slide in presentation.slides
            if slide.shapes[0].text.startswith("完整来源")
        )
    else:
        document_pdf = PdfReader(source)
        assert len(document_pdf.pages) > 1
        pages = []
        for number, page in enumerate(document_pdf.pages, 1):
            text = page.extract_text()
            label = "客户" if audience == "customer" else "内部"
            footer = f"{status} · {label} · {receipt['编号']} · {number}"
            assert footer in text
            pages.append(text.replace(footer, "", 1))
        extracted = "\n".join(pages)

    compact = "".join(extracted.split())
    assert "来源边界BEGIN" in compact and "来源边界END" in compact
    retained = "".join((source_text if source_text is not None else extracted).split())
    first = retained.index("来源边界BEGIN")
    last = retained.index("来源边界END") + len("来源边界END")
    assert retained[first:last] == "".join(original.split())
    assert "".join(formula.split()) in compact
    assert receipt["编号"] in extracted
    assert "未记录；不补零" in extracted
