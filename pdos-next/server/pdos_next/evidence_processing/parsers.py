"""Trusted parser implementations. The public entry point runs these in a bounded child."""

import hashlib
import io
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from xml.etree import ElementTree

from .types import ParseFailure, ParseLimits, ParseResult, SourceLocation, TextSegment

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SUPPORTED_MEDIA = frozenset({"text/plain", "text/markdown", "application/pdf", XLSX})


def _fail(code: str, message: str):
    raise ParseFailure(code, message)


class _Segments:
    def __init__(self, limits: ParseLimits):
        self.limits = limits
        self.values: list[TextSegment] = []
        self.chars = 0

    def add(self, text: str, location: SourceLocation, kind: str = "text"):
        if len(self.values) >= self.limits.max_segments:
            _fail("OUTPUT_LIMIT", "Original contains too many extractable source locations")
        self.chars += len(text) + 2
        if self.chars > self.limits.max_text_chars:
            _fail("OUTPUT_LIMIT", "Extracted text exceeds the configured character limit")
        self.values.append(TextSegment(text, location, kind))


def _utf8(content: bytes, limits: ParseLimits) -> tuple[str, tuple[TextSegment, ...]]:
    try:
        text = content.decode("utf-8", errors="strict")
    except UnicodeDecodeError as exc:
        raise ParseFailure("INVALID_ENCODING", "Text original must use UTF-8") from exc
    if "\x00" in text:
        _fail("MEDIA_TYPE_MISMATCH", "Text original contains binary null bytes")
    if not text.strip():
        _fail("EMPTY_CONTENT", "Text original has no nonempty content")
    if len(text) > limits.max_text_chars:
        _fail("OUTPUT_LIMIT", "Text exceeds the configured character limit")
    lines = text.splitlines()
    if len(lines) > limits.max_segments:
        _fail("OUTPUT_LIMIT", "Text contains too many source lines")
    return text, tuple(
        TextSegment(line, SourceLocation("line", line=index))
        for index, line in enumerate(lines, start=1)
    )


def _xlsx_archive(content: bytes, limits: ParseLimits):
    if not content.startswith(b"PK\x03\x04"):
        _fail("MEDIA_TYPE_MISMATCH", "Spreadsheet original must be an XLSX archive")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > limits.max_zip_entries:
                _fail("ARCHIVE_LIMIT", "Spreadsheet contains too many archive entries")
            expanded = 0
            names = set()
            content_types = None
            for member in members:
                name = member.filename
                if (
                    name in names
                    or name.startswith(("/", "\\"))
                    or "\\" in name
                    or ".." in name.split("/")
                ):
                    _fail("INVALID_ARCHIVE", "Spreadsheet archive has invalid member paths")
                names.add(name)
                if member.flag_bits & 1:
                    _fail("ENCRYPTED_DOCUMENT", "Encrypted spreadsheet archives are unsupported")
                expanded += member.file_size
                if (
                    expanded > limits.max_zip_expanded_bytes
                    or member.file_size > max(1, member.compress_size) * limits.max_zip_ratio
                ):
                    _fail("ARCHIVE_LIMIT", "Spreadsheet expansion exceeds the configured limit")
                lower = name.lower()
                if "vbaproject" in lower or lower.endswith((".vba", ".exe", ".dll", ".js")):
                    _fail(
                        "ACTIVE_CONTENT", "Macro or executable spreadsheet content is unsupported"
                    )
                if lower.startswith("xl/externallinks/"):
                    _fail("EXTERNAL_CONTENT", "Externally linked workbook data is unsupported")
                if lower.startswith(("xl/macrosheets/", "xl/activex/", "xl/embeddings/")):
                    _fail("ACTIVE_CONTENT", "Embedded active workbook content is unsupported")
                if lower.endswith((".xml", ".rels")):
                    value = archive.read(member)
                    upper = value.upper()
                    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
                        _fail("ACTIVE_CONTENT", "Document entity declarations are unsupported")
                    if name == "[Content_Types].xml":
                        content_types = value
            if "xl/workbook.xml" not in names or content_types is None:
                _fail("MEDIA_TYPE_MISMATCH", "Archive is not an XLSX workbook")
            types = ElementTree.fromstring(content_types)
            workbook_type = False
            for entry in types:
                kind = entry.attrib.get("ContentType", "")
                if any(
                    marker in kind.lower()
                    for marker in ("macroenabled", "macrosheet", "vba", "activex")
                ):
                    _fail("ACTIVE_CONTENT", "Macro-enabled workbooks are unsupported")
                if entry.attrib.get("PartName") == "/xl/workbook.xml":
                    workbook_type = kind == (
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
                    )
            if not workbook_type:
                _fail("MEDIA_TYPE_MISMATCH", "Workbook content type does not match XLSX")
    except (zipfile.BadZipFile, ElementTree.ParseError, RuntimeError, OSError) as exc:
        raise ParseFailure("MALFORMED_DOCUMENT", "Spreadsheet archive is malformed") from exc


def _xlsx(content: bytes, limits: ParseLimits):
    _xlsx_archive(content, limits)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ParseFailure(
            "PARSER_UNAVAILABLE", "Spreadsheet parser is unavailable", retryable=True
        ) from exc
    workbook = None
    values = _Segments(limits)
    warnings = []
    try:
        # Formulas are retained as original expressions, never evaluated as measurements.
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        if len(workbook.worksheets) > limits.max_sheets:
            _fail("DOCUMENT_LIMIT", "Workbook contains too many worksheets")
        for sheet in workbook.worksheets:
            if sheet.max_row and sheet.max_row > limits.max_rows_per_sheet:
                _fail("DOCUMENT_LIMIT", "Worksheet exceeds the configured row limit")
            if sheet.max_column and sheet.max_column > limits.max_columns_per_sheet:
                _fail("DOCUMENT_LIMIT", "Worksheet exceeds the configured column limit")
            if sheet.sheet_state != "visible":
                warnings.append(f"Hidden worksheet retained: {sheet.title}")
            # A malicious dimension must not hide actual cells from the reader.
            sheet.reset_dimensions()
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                if (
                    row_number > limits.max_rows_per_sheet
                    or len(row) > limits.max_columns_per_sheet
                ):
                    _fail("DOCUMENT_LIMIT", "Worksheet actual cells exceed the configured limits")
                for column_number, cell in enumerate(row, start=1):
                    value = cell.value
                    if value is None:
                        continue
                    if isinstance(value, (datetime, date, time)):
                        text, kind = value.isoformat(), "date"
                    elif isinstance(value, bool):
                        text, kind = "true" if value else "false", "boolean"
                    elif isinstance(value, (int, float, Decimal)):
                        text, kind = str(value), "number"
                    else:
                        text = str(value)
                        kind = (
                            "formula"
                            if cell.data_type == "f"
                            else "error"
                            if cell.data_type == "e"
                            else "text"
                        )
                    values.add(
                        text,
                        SourceLocation(
                            "cell",
                            sheet=sheet.title,
                            row=row_number,
                            column=column_number,
                            cell=cell.coordinate,
                        ),
                        kind,
                    )
        if not values.values:
            _fail("EMPTY_CONTENT", "Workbook has no extractable populated cells")
        if any(item.value_kind == "formula" for item in values.values):
            warnings.append("Formula expressions were retained without evaluation")
        return "\n".join(item.text for item in values.values), tuple(values.values), tuple(warnings)
    except ParseFailure:
        raise
    except Exception as exc:
        raise ParseFailure("MALFORMED_DOCUMENT", "Spreadsheet could not be parsed") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _pdf(content: bytes, limits: ParseLimits):
    if not content.startswith(b"%PDF-"):
        _fail("MEDIA_TYPE_MISMATCH", "PDF original must contain a PDF header")
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ParseFailure(
            "PARSER_UNAVAILABLE", "PDF parser is unavailable", retryable=True
        ) from exc
    values = _Segments(limits)
    warnings = []
    try:
        reader = PdfReader(io.BytesIO(content), strict=True)
        if reader.is_encrypted:
            _fail("ENCRYPTED_DOCUMENT", "Encrypted PDFs require an explicit separate workflow")
        if len(reader.pages) > limits.max_pdf_pages:
            _fail("DOCUMENT_LIMIT", "PDF contains too many pages")
        root = reader.trailer["/Root"]
        names = root.get("/Names")
        if names is not None:
            names = names.get_object()
        if any(key in root for key in ("/OpenAction", "/AA")) or (
            names and any(key in names for key in ("/JavaScript", "/EmbeddedFiles"))
        ):
            _fail("ACTIVE_CONTENT", "PDF actions or embedded files are unsupported")
        for number, page in enumerate(reader.pages, start=1):
            if "/AA" in page:
                _fail("ACTIVE_CONTENT", "PDF page actions are unsupported")
            for annotation in page.get("/Annots", []):
                annotation = annotation.get_object()
                action = annotation.get("/A")
                if action is not None:
                    action = action.get_object()
                if "/AA" in annotation or (
                    action
                    and action.get("/S") in {"/JavaScript", "/Launch", "/SubmitForm", "/ImportData"}
                ):
                    _fail("ACTIVE_CONTENT", "PDF annotation actions are unsupported")
            stream = page.get_contents()
            if stream is not None and len(stream.get_data()) > limits.max_zip_expanded_bytes:
                _fail("DOCUMENT_LIMIT", "PDF page stream exceeds the configured limit")
            text = page.extract_text() or ""
            if text.strip():
                values.add(text, SourceLocation("page", page=number))
            else:
                warnings.append(f"Page {number} has no extracted text; OCR was not performed")
        return (
            "\n\n".join(item.text for item in values.values),
            tuple(values.values),
            tuple(warnings),
            "parsed" if values.values else "needs_ocr",
            tuple(
                SourceLocation("page", page=number) for number in range(1, len(reader.pages) + 1)
            ),
        )
    except ParseFailure:
        raise
    except Exception as exc:
        raise ParseFailure("MALFORMED_DOCUMENT", "PDF could not be parsed") from exc


def parse_in_worker(content: bytes, media_type: str, limits: ParseLimits) -> ParseResult:
    if media_type not in SUPPORTED_MEDIA:
        _fail("UNSUPPORTED_MEDIA_TYPE", "Original media type is not supported")
    if not content:
        _fail("EMPTY_CONTENT", "Original byte sequence is empty")
    if len(content) > limits.max_input_bytes:
        _fail("FILE_TOO_LARGE", "Original exceeds the configured byte limit")
    warnings, status = (), "parsed"
    expected = None
    if media_type in {"text/plain", "text/markdown"}:
        if content.startswith((b"%PDF-", b"PK\x03\x04")):
            _fail("MEDIA_TYPE_MISMATCH", "Declared text type does not match the original bytes")
        text, segments = _utf8(content, limits)
    elif media_type == XLSX:
        text, segments, warnings = _xlsx(content, limits)
    else:
        text, segments, warnings, status, expected = _pdf(content, limits)
    if expected is None:
        expected = tuple(dict.fromkeys(segment.location for segment in segments))
    return ParseResult(
        hashlib.sha256(content).hexdigest(),
        media_type,
        text,
        segments,
        status=status,
        warnings=warnings,
        expected_locations=expected,
    )
